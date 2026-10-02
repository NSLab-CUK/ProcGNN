"""Generate only Fig. 4a from the current manuscript's single-process table.

This script reads existing Table 2 values from experiment_results.xlsx. It
does not train, fine-tune, or evaluate any model.
"""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from matplotlib.lines import Line2D


ROOT = Path(__file__).resolve().parents[1]
SOURCE_WORKBOOK = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "experiment_results"
    / "experiment_results.xlsx"
)
SOURCE_SHEET = "Table2_Single_10D"
OUT_DIR = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "Fig4a_single_process_performance"
)

# Source heading, requested manuscript display name, publication tick label.
PROPERTY_SPECS: tuple[tuple[str, str, str], ...] = (
    ("H₂O", "H2O", r"H$_2$O"),
    ("H₂", "H2", r"H$_2$"),
    ("CH₄", "CH4", r"CH$_4$"),
    ("CO₂", "CO2", r"CO$_2$"),
    ("CO", "CO", "CO"),
    ("O₂", "O2", r"O$_2$"),
    ("N₂", "N2", r"N$_2$"),
    ("Temp", "Temperature", "Temperature"),
    ("Pres", "Pressure", "Pressure"),
    ("Mass Flow", "Mass Flow", "Mass Flow"),
)
NUMBER_RE = re.compile(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?")


def relative_path(path: Path) -> str:
    """Print source/output paths relative to the repository root."""
    return path.resolve().relative_to(ROOT.resolve()).as_posix()


def parse_mean_and_sd(cell_value: object, *, context: str) -> tuple[float, float] | tuple[float, float]:
    """Parse the existing manuscript's ``mean ± SD`` cell without rounding it."""
    if cell_value is None or not str(cell_value).strip():
        return math.nan, math.nan
    values = [float(token) for token in NUMBER_RE.findall(str(cell_value))]
    if len(values) < 2:
        raise ValueError(f"{context}: expected a mean and SD in {cell_value!r}")
    return values[0], values[1]


def source_column_map(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[tuple[str, str], int]:
    """Recover merged Table 2 headers and map (source property, metric) to Excel columns."""
    map_: dict[tuple[str, str], int] = {}
    current_property: str | None = None
    for column in range(1, ws.max_column + 1):
        property_header = ws.cell(row=1, column=column).value
        if property_header is not None:
            current_property = str(property_header)
        metric = ws.cell(row=2, column=column).value
        if current_property is not None and metric is not None:
            map_[(current_property, str(metric))] = column
    return map_


def extract_source_table() -> tuple[pd.DataFrame, list[str]]:
    """Extract all Table 2 model/property MAE and sMAPE values exactly once."""
    if not SOURCE_WORKBOOK.is_file():
        raise FileNotFoundError(f"Current manuscript workbook is missing: {SOURCE_WORKBOOK}")
    workbook = openpyxl.load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise KeyError(f"Workbook has no {SOURCE_SHEET!r} sheet: {workbook.sheetnames}")
    worksheet = workbook[SOURCE_SHEET]
    column_map = source_column_map(worksheet)

    source_properties = [source_name for source_name, _, _ in PROPERTY_SPECS]
    available_properties = {
        property_name
        for property_name, metric in column_map
        if metric in {"MAE", "sMAPE (%)"} and property_name != "Average"
    }
    if set(source_properties) != available_properties:
        raise ValueError(
            "Table 2 target-property headings do not match the requested figure scope; "
            f"expected={source_properties}, available={sorted(available_properties)}"
        )
    for property_name in source_properties:
        for metric in ("MAE", "sMAPE (%)"):
            if (property_name, metric) not in column_map:
                raise ValueError(f"Table 2 is missing {metric} for {property_name}")

    records: list[dict[str, object]] = []
    # Table 2 model rows begin at row 3 and end at the first completely blank model cell.
    for row_index in range(3, worksheet.max_row + 1):
        model_cell = worksheet.cell(row=row_index, column=1).value
        if model_cell is None or not str(model_cell).strip():
            break
        model = str(model_cell).strip()
        for source_name, property_name, _ in PROPERTY_SPECS:
            mae_cell = worksheet.cell(row=row_index, column=column_map[(source_name, "MAE")])
            smape_cell = worksheet.cell(row=row_index, column=column_map[(source_name, "sMAPE (%)")])
            mae_mean, mae_sd = parse_mean_and_sd(
                mae_cell.value,
                context=f"{model} / {source_name} / MAE",
            )
            smape_mean, smape_sd = parse_mean_and_sd(
                smape_cell.value,
                context=f"{model} / {source_name} / sMAPE",
            )
            records.append(
                {
                    "Model": model,
                    "Property": property_name,
                    "Source property": source_name,
                    "MAE": mae_mean,
                    "MAE SD": mae_sd,
                    "sMAPE (%)": smape_mean,
                    "sMAPE SD (%)": smape_sd,
                    "Source MAE cell": str(mae_cell.value or ""),
                    "Source sMAPE cell": str(smape_cell.value or ""),
                    "Source workbook": relative_path(SOURCE_WORKBOOK),
                    "Source worksheet": SOURCE_SHEET,
                }
            )
    workbook.close()

    source = pd.DataFrame(records)
    if source.empty:
        raise ValueError("No model/property rows were extracted from Table 2.")
    expected_rows = source["Model"].nunique() * len(PROPERTY_SPECS)
    if len(source) != expected_rows or source.duplicated(["Model", "Property"]).any():
        raise ValueError("Table 2 extraction did not yield one MAE/sMAPE record per model and property.")
    return source, source_properties


def make_comparison(source: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Select the minimum existing non-proposed baseline for each property."""
    models = source["Model"].drop_duplicates().tolist()
    if models.count("Proposed") != 1:
        raise ValueError(f"Expected exactly one Proposed row in Table 2; found {models.count('Proposed')}")

    source = source.copy()
    source["Record type"] = "source_metric"
    # Keep the workbook's percent values for provenance while storing the
    # unitless ratio used on the figure's sMAPE axis (e.g., 20% -> 0.20).
    source["sMAPE (ratio)"] = source["sMAPE (%)"] / 100.0
    source["Comparison eligibility"] = np.where(
        source["Model"].eq("Proposed"),
        "proposed",
        np.where(source["Model"].eq("Proposed (All edge)"), "excluded_all_edge_diagnostic", "baseline"),
    )

    comparisons: list[dict[str, object]] = []
    for source_name, property_name, _ in PROPERTY_SPECS:
        property_rows = source.loc[source["Property"].eq(property_name)].copy()
        proposed_rows = property_rows.loc[property_rows["Model"].eq("Proposed")]
        if len(proposed_rows) != 1 or not np.isfinite(proposed_rows.iloc[0]["sMAPE (%)"]):
            raise ValueError(f"{property_name}: unavailable or non-unique Proposed sMAPE in Table 2")
        proposed_smape = float(proposed_rows.iloc[0]["sMAPE (%)"])

        baselines = property_rows.loc[
            property_rows["Comparison eligibility"].eq("baseline")
            & np.isfinite(property_rows["sMAPE (%)"])
        ].copy()
        if baselines.empty:
            raise ValueError(f"{property_name}: no finite baseline sMAPE values in Table 2")
        best_smape = float(baselines["sMAPE (%)"].min())
        tied_names = baselines.loc[baselines["sMAPE (%)"].eq(best_smape), "Model"].tolist()
        relative_reduction = 100.0 * (best_smape - proposed_smape) / best_smape
        comparisons.append(
            {
                "Record type": "plot_comparison",
                "Model": "",
                "Property": property_name,
                "Source property": source_name,
                "MAE": math.nan,
                "MAE SD": math.nan,
                "sMAPE (%)": math.nan,
                "sMAPE SD (%)": math.nan,
                "sMAPE (ratio)": math.nan,
                "Source MAE cell": "",
                "Source sMAPE cell": "",
                "Source workbook": relative_path(SOURCE_WORKBOOK),
                "Source worksheet": SOURCE_SHEET,
                "Comparison eligibility": "plot_pair",
                "Proposed sMAPE (%)": proposed_smape,
                "Proposed sMAPE (ratio)": proposed_smape / 100.0,
                "Best baseline": "; ".join(tied_names),
                "Best baseline tied names": "; ".join(tied_names),
                "Best baseline sMAPE (%)": best_smape,
                "Best baseline sMAPE (ratio)": best_smape / 100.0,
                "Relative reduction (%)": relative_reduction,
                "Proposed is best": proposed_smape <= best_smape,
            }
        )
    comparison = pd.DataFrame(comparisons)
    return source, comparison


def output_table(source: pd.DataFrame, comparison: pd.DataFrame) -> pd.DataFrame:
    """Store all source metrics plus the exact numerical pairs used in the plot."""
    fields = [
        "Record type",
        "Model",
        "Property",
        "Source property",
        "MAE",
        "MAE SD",
        "sMAPE (%)",
        "sMAPE SD (%)",
        "sMAPE (ratio)",
        "Comparison eligibility",
        "Proposed sMAPE (%)",
        "Proposed sMAPE (ratio)",
        "Best baseline",
        "Best baseline tied names",
        "Best baseline sMAPE (%)",
        "Best baseline sMAPE (ratio)",
        "Relative reduction (%)",
        "Proposed is best",
        "Source MAE cell",
        "Source sMAPE cell",
        "Source workbook",
        "Source worksheet",
    ]
    for field in fields:
        if field not in source:
            source[field] = pd.NA
        if field not in comparison:
            comparison[field] = pd.NA
    # Object dtype retains blank comparison-only/source-only cells without
    # coercing an all-missing column during concatenation.
    source_output = source.reindex(columns=fields).astype(object)
    comparison_output = comparison.reindex(columns=fields).astype(object)
    combined = pd.concat([source_output, comparison_output], ignore_index=True)
    return combined


def inspect_scale(comparison: pd.DataFrame) -> str:
    values = comparison[["Proposed sMAPE (ratio)", "Best baseline sMAPE (ratio)"]].to_numpy(float).ravel()
    if not np.all(np.isfinite(values)) or np.any(values <= 0):
        raise ValueError("sMAPE values must be finite and positive before selecting a log scale.")
    ratio = float(values.max() / values.min())
    # The existing values span more than two orders of magnitude. A log axis
    # preserves the exact values while keeping the N2 pair visible.
    if ratio > 100.0:
        print(f"Scale inspection: sMAPE range={values.min():.6f}-{values.max():.6f} (ratio={ratio:.1f}); using log scale.")
        return "log"
    print(f"Scale inspection: sMAPE range={values.min():.6f}-{values.max():.6f} (ratio={ratio:.1f}); using linear scale.")
    return "linear"


def draw_figure(
    comparison: pd.DataFrame,
    pdf_path: Path,
    png_path: Path,
    *,
    title: str = "(a) Single-Process Prediction",
) -> tuple[float, float]:
    figure_size = (7.6, 5.05)
    tick_labels = [label for _, _, label in PROPERTY_SPECS]
    x_positions = np.arange(len(comparison), dtype=float)
    baseline = comparison["Best baseline sMAPE (ratio)"].to_numpy(float)
    proposed = comparison["Proposed sMAPE (ratio)"].to_numpy(float)
    scale = inspect_scale(comparison)

    with plt.rc_context(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10.5,
            "axes.titlesize": 15,
            "axes.labelsize": 11,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "figure.facecolor": "white",
            "savefig.facecolor": "white",
        }
    ):
        fig, ax = plt.subplots(figsize=figure_size, layout="constrained")
        offset = 0.105
        for x_value, baseline_value, proposed_value in zip(x_positions, baseline, proposed):
            ax.plot(
                [x_value - offset, x_value + offset],
                [baseline_value, proposed_value],
                color="#B9C2CC",
                linewidth=1.35,
                zorder=1,
            )
        ax.scatter(
            x_positions - offset,
            baseline,
            s=48,
            marker="o",
            color="#6D7F92",
            edgecolor="white",
            linewidth=0.8,
            zorder=3,
        )
        ax.scatter(
            x_positions + offset,
            proposed,
            s=55,
            marker="D",
            color="#2166AC",
            edgecolor="white",
            linewidth=0.9,
            zorder=4,
        )

        if scale == "log":
            ax.set_yscale("log")
            ax.set_ylim(0.0005, 0.40)
            ax.set_yticks([0.001, 0.003, 0.01, 0.03, 0.1, 0.3])
            ax.set_yticklabels(["0.001", "0.003", "0.01", "0.03", "0.1", "0.3"])
            y_label = "sMAPE (log scale)"
            legend_floor = 0.15
        else:
            upper = float(max(baseline.max(), proposed.max()) * 1.15)
            ax.set_ylim(bottom=0.0, top=upper)
            y_label = "sMAPE"
            legend_floor = upper * 0.50

        ax.set_xlim(-0.55, len(comparison) - 0.45)
        ax.set_xticks(x_positions, tick_labels)
        # The final three manuscript property names are longer than chemical
        # formulas. Rotate only those labels enough to keep them distinct.
        for tick_label in ax.get_xticklabels()[-3:]:
            tick_label.set_rotation(30)
            tick_label.set_ha("right")
            tick_label.set_rotation_mode("anchor")
        ax.set_ylabel(y_label, labelpad=9)
        ax.set_title(title, pad=18, fontweight="normal", color="#20252B")
        ax.grid(axis="y", which="major", color="#E3E8EE", linewidth=0.7)
        ax.grid(axis="y", which="minor", visible=False)
        ax.set_axisbelow(True)
        ax.tick_params(axis="x", length=0, pad=7)
        ax.tick_params(axis="y", length=3.5, width=0.8)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.spines["left"].set_color("#7D8A99")
        ax.spines["bottom"].set_color("#7D8A99")

        # The selected in-axis legend must sit above every right-side marker.
        right_side = comparison.iloc[7:][["Proposed sMAPE (ratio)", "Best baseline sMAPE (ratio)"]].to_numpy(float)
        if right_side.max() >= legend_floor:
            raise ValueError("The selected upper-right legend region is no longer empty.")
        legend = ax.legend(
            handles=[
                Line2D(
                    [0], [0], marker="D", color="none", markerfacecolor="#2166AC",
                    markeredgecolor="white", markersize=7.5, label="Proposed"
                ),
                Line2D(
                    [0], [0], marker="o", color="none", markerfacecolor="#6D7F92",
                    markeredgecolor="white", markersize=7.2, label="Best baseline"
                ),
            ],
            loc="upper right",
            bbox_to_anchor=(0.985, 0.985),
            borderaxespad=0.25,
            frameon=True,
            framealpha=0.98,
            facecolor="white",
            edgecolor="#CFD7DF",
            fontsize=9.4,
            labelspacing=0.45,
            handletextpad=0.55,
        )
        legend.get_frame().set_linewidth(0.7)

        fig.savefig(pdf_path, format="pdf", bbox_inches="tight", pad_inches=0.05)
        fig.savefig(png_path, format="png", dpi=600, bbox_inches="tight", pad_inches=0.05)
        plt.close(fig)
    return figure_size


def prepare_outputs(overwrite: bool) -> dict[str, Path]:
    outputs = {
        "pdf": OUT_DIR / "Fig4a_single_process_performance.pdf",
        "png": OUT_DIR / "Fig4a_single_process_performance_600dpi.png",
        "csv": OUT_DIR / "Fig4a_single_process_performance.csv",
    }
    existing = [path for path in outputs.values() if path.exists()]
    if existing and not overwrite:
        names = ", ".join(path.name for path in existing)
        raise FileExistsError(f"Refusing to overwrite existing Fig. 4a outputs: {names}. Use --overwrite.")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if overwrite:
        for path in existing:
            path.unlink()
    return outputs


def print_validation(source: pd.DataFrame, comparison: pd.DataFrame, source_properties: list[str]) -> None:
    print("\nTarget-property headings in current Table 2:")
    print("- " + ", ".join(source_properties))
    print(f"Model rows extracted: {source['Model'].nunique()} (baseline candidates: {(source['Comparison eligibility'].eq('baseline')).groupby(source['Model']).first().sum()})")
    print("\nProperty comparison (sMAPE %):")
    display = comparison.loc[
        :,
        [
            "Property",
            "Proposed sMAPE (%)",
            "Best baseline",
            "Best baseline sMAPE (%)",
            "Relative reduction (%)",
        ],
    ].copy()
    print(display.to_string(index=False, float_format=lambda value: f"{value:.4f}"))

    proposed_best_count = int(comparison["Proposed is best"].sum())
    not_best = comparison.loc[~comparison["Proposed is best"], "Property"].tolist()
    proposed_mean = float(comparison["Proposed sMAPE (%)"].mean())
    baseline_mean = float(comparison["Best baseline sMAPE (%)"].mean())
    print(f"\nProperties where Proposed is best: {proposed_best_count}/{len(comparison)}")
    print(f"Properties where Proposed is not best: {', '.join(not_best) if not_best else 'None'}")
    print(f"Mean sMAPE of Proposed: {proposed_mean:.4f}%")
    print(f"Mean sMAPE of property-wise best baseline: {baseline_mean:.4f}%")
    print(
        "WARNING: The current Table 2 caption does not state an independent count of properties "
        "where Proposed is best; the displayed sMAPE means yield the reported count above."
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--overwrite", action="store_true", help="replace only the three Fig. 4a outputs")
    args = parser.parse_args()

    print("Source files used:")
    print(f"- {relative_path(SOURCE_WORKBOOK)} (worksheet: {SOURCE_SHEET})")
    source, source_properties = extract_source_table()
    source, comparison = make_comparison(source)
    outputs = prepare_outputs(args.overwrite)
    output_table(source, comparison).to_csv(outputs["csv"], index=False, float_format="%.10g")
    figure_size = draw_figure(comparison, outputs["pdf"], outputs["png"])
    print_validation(source, comparison, source_properties)
    print(f"\nFigure size: {figure_size[0]:.2f} x {figure_size[1]:.2f} inches")
    print("Output paths:")
    for path in outputs.values():
        print(f"- {relative_path(path)}")


if __name__ == "__main__":
    main()
