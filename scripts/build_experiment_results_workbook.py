"""Build one traceable workbook for the latest paper-result datasets.

The workbook intentionally combines only the result families used by the final
paper tables/figures.  It keeps paper-ready comparison tables in their
original layout, while long-format figure data remain numeric and filterable.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from copy import copy
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "experiment_results"
    / "experiment_results.xlsx"
)

PAPER_TABLES = ROOT / "outputs" / "0819final" / "Paper_Tables_Final3.xlsx"
TARGET_WORKBOOK = (
    ROOT
    / "outputs"
    / "0819final"
    / "Final3_target_performance"
    / "Final3_target_only_tables.xlsx"
)
FINAL_DATA = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "final_model_results"
    / "data"
)
SENSITIVITY_CSV = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "sensitivity_7cases_center3_measured_outer4_taller_font_1p7x_v2_20260917"
    / "sensitivity_7cases_data.csv"
)
CONSTRAINT_WORKBOOK = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "constraint_satisfaction_ablation_grouped_bars_typography_20260921_v4_top_left_legend"
    / "constraint_satisfaction_grouped_bar_data.xlsx"
)
TARGET_SHAP_ROOT = (
    ROOT
    / "outputs"
    / "0819final"
    / "Paper_Tables_Final3_figures"
    / "final"
    / "shap_target_material_27_flowsheets_20260922"
)

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
GRAY = "F2F2F2"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="D9E2F3")


def relative(path: Path) -> str:
    return path.resolve().relative_to(ROOT.resolve()).as_posix()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require_file(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"Required source file is missing: {path}")


def safe_title(value: str) -> str:
    return value[:31]


def parse_csv_value(value: str):
    """Keep categorical strings intact while making numeric CSV cells usable in Excel."""
    stripped = value.strip()
    if stripped == "":
        return None
    try:
        numeric = float(stripped)
    except ValueError:
        return value
    if numeric.is_integer() and all(token not in stripped.lower() for token in (".", "e")):
        return int(numeric)
    return numeric


def style_sheet(ws, *, freeze: str = "A2", autofilter: bool = True) -> None:
    ws.freeze_panes = freeze
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.outlinePr.summaryBelow = False
    if autofilter and ws.max_row >= 2 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = copy(cell.alignment) if cell.has_style else Alignment(vertical="center")
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal or "left",
                vertical="center",
                wrap_text=True,
            )
    for column in range(1, ws.max_column + 1):
        width = 12
        for row in range(1, min(ws.max_row, 300) + 1):
            value = ws.cell(row=row, column=column).value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 42))
        ws.column_dimensions[get_column_letter(column)].width = width


def add_csv_sheet(book: Workbook, title: str, path: Path, *, notes: str | None = None) -> None:
    require_file(path)
    ws = book.create_sheet(safe_title(title))
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    if not rows:
        raise ValueError(f"CSV source has no rows: {path}")
    for row in rows:
        ws.append([parse_csv_value(value) for value in row])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    style_sheet(ws)
    if notes:
        ws.sheet_properties.tabColor = "5B9BD5"


def copy_source_sheet(
    book: Workbook,
    source_path: Path,
    source_name: str,
    target_name: str,
    *,
    freeze: str,
    filter_row: int | None = None,
) -> None:
    require_file(source_path)
    source_book = load_workbook(source_path, data_only=True, read_only=False)
    source = source_book[source_name]
    target = book.create_sheet(safe_title(target_name))

    for row in source.iter_rows():
        for source_cell in row:
            cell = target.cell(row=source_cell.row, column=source_cell.column, value=source_cell.value)
            if source_cell.has_style:
                cell._style = copy(source_cell._style)
                cell.number_format = source_cell.number_format
                cell.font = copy(source_cell.font)
                cell.fill = copy(source_cell.fill)
                cell.border = copy(source_cell.border)
                cell.alignment = copy(source_cell.alignment)
            if source_cell.hyperlink:
                cell._hyperlink = copy(source_cell.hyperlink)

    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
        target.column_dimensions[key].hidden = dimension.hidden
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key].height = dimension.height
        target.row_dimensions[key].hidden = dimension.hidden

    target.freeze_panes = freeze
    target.sheet_view.showGridLines = False
    target.sheet_properties.pageSetUpPr.fitToPage = True
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 0
    if filter_row is not None:
        target.auto_filter.ref = (
            f"A{filter_row}:{get_column_letter(target.max_column)}{target.max_row}"
        )
    source_book.close()


def make_readme(book: Workbook, output: Path) -> None:
    ws = book.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 115
    ws.merge_cells("A1:B1")
    ws["A1"] = "Latest Experiment Results"
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    rows = [
        ("Generated", date.today().isoformat()),
        ("Purpose", "One traceable workbook for the latest paper tables and figure values."),
        ("Scope", "Single-/multi-process comparison, module and physics ablations, transfer learning, computational efficiency, sensitivity, conservation satisfaction, absolute SHAP, and 27 target-edge-material SHAP results."),
        ("Primary metric", "Table 2/3 and ablation sheets retain displayed MAE, NMAE, and sMAPE values. Lower is better for all three error metrics."),
        ("Target scope", "Target-edge performance is retained separately in Target_Performance_Long; no validation/checkpoint monitor keys are included."),
        ("Efficiency caution", "Target-stream timing/performance values are benchmarked. All-stream baseline values are documented scenario extensions; provenance columns must be checked before making matched-benchmark claims."),
        ("Sensitivity caution", "Sensitivity_7Case carries the source-level Data provenance and Uncertainty provenance columns. Rows marked illustrative are not measured results."),
        ("SHAP", "Absolute SHAP values use the final model package. Target_SHAP_27 contains signed and absolute relative SHAP summaries for 10 H2, 10 CO2, and 7 H2O target-edge-material explanations."),
        ("Output", relative(output)),
        ("Sheet guide", "Tables 2/3: comparisons; Ablation_*: architecture and PINN objectives; Transfer_sMAPE/Efficiency/Sensitivity/Constraint_*: figure data; SHAP_*: final attribution values; Target_Performance_Long: numeric detailed source table."),
    ]
    for row_index, (label, value) in enumerate(rows, start=3):
        ws.cell(row=row_index, column=1, value=label)
        ws.cell(row=row_index, column=2, value=value)
        ws.cell(row=row_index, column=1).font = Font(bold=True, color=NAVY)
        ws.cell(row=row_index, column=1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row=row_index, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_index, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_index].height = 32 if row_index in {5, 8, 9, 11, 12} else 22
    ws.freeze_panes = "A3"


def add_source_index(book: Workbook, records: Iterable[tuple[str, Path, str, str]]) -> None:
    ws = book.create_sheet("Sources")
    ws.append(["Result family", "Source file", "SHA-256", "Contents", "Status / caution"])
    for family, path, contents, caution in records:
        require_file(path)
        ws.append([family, relative(path), sha256(path), contents, caution])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in range(2, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=column)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
        ws.row_dimensions[row].height = 42
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 66
    ws.column_dimensions["D"].width = 47
    ws.column_dimensions["E"].width = 70
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"
    ws.sheet_view.showGridLines = False


def decorate_table_sheet(book: Workbook, title: str, *, header_row: int = 1) -> None:
    ws = book[title]
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color=WHITE, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    ws.sheet_properties.tabColor = "5B9BD5"


def build(output: Path) -> None:
    if output.exists():
        raise FileExistsError(
            f"Refusing to overwrite existing workbook: {output}. Choose a new path or remove it explicitly."
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    make_readme(book, output)

    sources = [
        (
            "Paper comparison and ablation tables",
            PAPER_TABLES,
            "Table 2/3 and module/PINN ablation performance plus conservation-satisfaction tables.",
            "Paper-ready displayed values; retain source formatting and metric units.",
        ),
        (
            "Detailed target-edge performance",
            TARGET_WORKBOOK,
            "Numeric long-format target-edge performance values extracted for final comparisons.",
            "Target edges only; use the metric/property columns rather than averaging incompatible units.",
        ),
        (
            "Transfer-learning figure data",
            FINAL_DATA / "transfer_learning" / "plot_data.csv",
            "Latest 0-10% and 0-100% transfer values for all models and ten properties.",
            "Numeric plotting data; sMAPE fields are percentages.",
        ),
        (
            "Computational efficiency figure data",
            FINAL_DATA / "efficiency" / "plot_data.csv",
            "Training/inference time and target/all-stream sMAPE data.",
            "All-stream baseline entries are documented scenario extensions; inspect provenance columns.",
        ),
        (
            "Seven-case sensitivity",
            SENSITIVITY_CSV,
            "Depth and mass/component/atom conservation-weight values.",
            "The included provenance fields distinguish supplied illustrative values from measured ones.",
        ),
        (
            "Conservation-satisfaction ablation",
            CONSTRAINT_WORKBOOK,
            "Module and physics-objective satisfaction-rate values for 1%, 5%, and 10% tolerances.",
            "Mean and standard deviation are kept as numeric values.",
        ),
        (
            "Absolute SHAP",
            FINAL_DATA / "absolute_shap" / "all_panels.csv",
            "Final-model mean absolute relative SHAP by unit type, operating variable, and feed type.",
            "Final package values; standard deviations are across processes.",
        ),
        (
            "Target-edge-material SHAP",
            TARGET_SHAP_ROOT / "shap_target_edge_node_summary.csv",
            "Node-level signed and absolute relative SHAP summaries for 27 selected target-edge-material results.",
            "Final-checkpoint target-edge explanations: H2=10, CO2=10, H2O=7.",
        ),
        (
            "Target-edge-material map",
            TARGET_SHAP_ROOT / "target_edge_species_map.csv",
            "Process, target edge, species, and target property mapping for the 27 explanations.",
            "Pairs the SHAP rows to their target edge/material.",
        ),
    ]
    add_source_index(book, sources)

    paper_sheets = [
        ("singleprocess_10D", "Table2_Single_10D", "B3", 2),
        ("singleprocess_physics", "Table2_Single_Phys", "B3", 2),
        ("multiprocess_10D", "Table3_Multi_10D", "B3", 2),
        ("multiprocess_physics", "Table3_Multi_Phys", "B3", 2),
        ("ablation_10D", "Ablation_Module_10D", "F3", 2),
        ("ablation_physics", "Ablation_Module_Phys", "F3", 2),
        ("pinn_ablation_10D", "Ablation_PINN_10D", "E3", 2),
        ("pinn_ablation_physics", "Ablation_PINN_Phys", "E3", 2),
    ]
    for source_name, target_name, freeze, filter_row in paper_sheets:
        copy_source_sheet(
            book,
            PAPER_TABLES,
            source_name,
            target_name,
            freeze=freeze,
            filter_row=filter_row,
        )
        book[target_name].sheet_properties.tabColor = "4472C4"

    add_csv_sheet(book, "Transfer_sMAPE", FINAL_DATA / "transfer_learning" / "plot_data.csv")
    add_csv_sheet(book, "Efficiency", FINAL_DATA / "efficiency" / "plot_data.csv")
    add_csv_sheet(book, "Sensitivity_7Case", SENSITIVITY_CSV)
    add_csv_sheet(book, "SHAP_Unit", FINAL_DATA / "absolute_shap" / "unit_type.csv")
    add_csv_sheet(book, "SHAP_Operating", FINAL_DATA / "absolute_shap" / "operating_variable.csv")
    add_csv_sheet(book, "SHAP_Feed", FINAL_DATA / "absolute_shap" / "feed_type.csv")
    add_csv_sheet(book, "Target_SHAP_27", TARGET_SHAP_ROOT / "shap_target_edge_node_summary.csv")
    add_csv_sheet(book, "Target_SHAP_Map", TARGET_SHAP_ROOT / "target_edge_species_map.csv")

    copy_source_sheet(
        book,
        CONSTRAINT_WORKBOOK,
        "Module ablation",
        "Constraint_Module",
        freeze="A2",
        filter_row=1,
    )
    copy_source_sheet(
        book,
        CONSTRAINT_WORKBOOK,
        "Physics objective",
        "Constraint_PINN",
        freeze="A2",
        filter_row=1,
    )
    for title in ("Constraint_Module", "Constraint_PINN"):
        decorate_table_sheet(book, title)
        style_sheet(book[title])

    copy_source_sheet(
        book,
        TARGET_WORKBOOK,
        "All_Target_Long",
        "Target_Performance_Long",
        freeze="A3",
        filter_row=2,
    )
    ws = book["Target_Performance_Long"]
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30
    ws.sheet_properties.tabColor = "70AD47"

    for name in book.sheetnames:
        ws = book[name]
        if name not in {"README", "Sources"} and ws.max_row > 0:
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0

    book.save(output)
    book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output if args.output.is_absolute() else (ROOT / args.output)
    build(output.resolve())
    print(f"Created {relative(output.resolve())}")


if __name__ == "__main__":
    main()
