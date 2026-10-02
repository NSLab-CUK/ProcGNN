#!/usr/bin/env python3
"""Add completed physics-conservation metrics to Final2_checked.xlsx.

The script reads the existing cross-server evaluation CSV only.  It does not
rerun a model or revise predictive-performance values.
"""
from __future__ import annotations

import argparse
import os
from copy import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


METRICS = [
    "normalized_abs_residual_mae",
    "normalized_residual_rmse",
    "normalized_abs_residual_max",
    "raw_huber_loss",
    "satisfaction_rate_abs_le_0.01",
    "satisfaction_rate_abs_le_0.05",
    "satisfaction_rate_abs_le_0.10",
]
DISPLAY = {
    "normalized_abs_residual_mae": "Normalized |residual| MAE\n(lower is better)",
    "normalized_residual_rmse": "Normalized residual RMSE\n(lower is better)",
    "normalized_abs_residual_max": "Maximum normalized |residual|\n(lower is better)",
    "raw_huber_loss": "Raw Huber loss\n(lower is better)",
    "satisfaction_rate_abs_le_0.01": "Pass rate |r| <= 1%\n(higher is better)",
    "satisfaction_rate_abs_le_0.05": "Pass rate |r| <= 5%\n(higher is better)",
    "satisfaction_rate_abs_le_0.10": "Pass rate |r| <= 10%\n(higher is better)",
}
PHASE_LABELS = {
    "Phase 1": "Joint multi-process",
    "Phase 1-S": "Single-process",
}
TERM_ORDER = {"mass": 0, "component": 1, "atom": 2}


def _mean_sd(values: pd.Series) -> tuple[float, float]:
    numeric = pd.to_numeric(values, errors="coerce").dropna()
    return float(numeric.mean()), float(numeric.std(ddof=1)) if len(numeric) > 1 else 0.0


def _mean_sd_text(values: pd.Series, percent: bool = False) -> str:
    mean, sd = _mean_sd(values)
    if percent:
        return f"{100.0 * mean:.4f}% +/- {100.0 * sd:.4f}%"
    return f"{mean:.4f} +/- {sd:.4f}"


def _build_summary(source: Path) -> pd.DataFrame:
    raw = pd.read_csv(source)
    required = {"phase", "family", "model", "condition", "term", *METRICS}
    missing = sorted(required - set(raw.columns))
    if missing:
        raise RuntimeError(f"physics source lacks columns: {missing}")
    raw = raw[(raw["model"].eq("Proposed")) & raw["phase"].isin(PHASE_LABELS)].copy()
    rows: list[dict[str, object]] = []
    for (phase, family, term), group in raw.groupby(["phase", "family", "term"], sort=False):
        row: dict[str, object] = {
            "Evaluation scope": PHASE_LABELS[str(phase)],
            "Source phase": phase,
            "Physics term": str(term).title(),
            "Logical runs": int(len(group)),
            "Graphs evaluated (total)": int(pd.to_numeric(group["graphs"], errors="coerce").fillna(0).sum()),
            "Valid residual entries (total)": int(pd.to_numeric(group["valid_residual_count"], errors="coerce").fillna(0).sum()),
        }
        for metric in METRICS:
            row[DISPLAY[metric].replace("\n", " ")] = _mean_sd_text(
                group[metric], percent=metric.startswith("satisfaction_rate")
            )
        rows.append(row)
    summary = pd.DataFrame(rows)
    if summary.empty:
        raise RuntimeError("no completed Proposed physics-conservation rows found")
    summary["_phase_order"] = summary["Source phase"].map({"Phase 1": 0, "Phase 1-S": 1})
    summary["_term_order"] = summary["Physics term"].str.lower().map(TERM_ORDER)
    return summary.sort_values(["_phase_order", "_term_order"]).drop(columns=["_phase_order", "_term_order"])


def _write_summary_sheet(book, summary: pd.DataFrame) -> None:
    if "Physics_Conservation" in book.sheetnames:
        del book["Physics_Conservation"]
    sheet = book.create_sheet("Physics_Conservation")
    headers = list(summary.columns)
    sheet.append(["Completed physics-conservation evaluation (target-edge model outputs)"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = sheet.cell(1, 1)
    title.font = Font(bold=True, color="FFFFFF", size=12)
    title.fill = PatternFill("solid", fgColor="1F4E78")
    title.alignment = Alignment(horizontal="left")
    sheet.append(headers)
    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for values in summary.itertuples(index=False, name=None):
        sheet.append(list(values))
    note_row = sheet.max_row + 2
    note = (
        "All values are mean +/- sample SD over logical runs, except totals. "
        "Residual terms are absolute normalized relative node-balance residuals. "
        "Pass rates are fractions satisfying |r| at the stated threshold."
    )
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(headers))
    note_cell = sheet.cell(note_row, 1, note)
    note_cell.font = Font(italic=True, color="666666", size=9)
    note_cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 34
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{sheet.max_row - 2}"
    widths = [24, 14, 16, 14, 23, 29, 27, 29, 30, 20, 22, 22, 23]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 42
    sheet.sheet_view.showGridLines = False


def _clarify_legacy_pass_rate_headers(book) -> None:
    for name in ("Single_Process", "Multi_Process"):
        if name not in book.sheetnames:
            continue
        sheet = book[name]
        # B1:D1 is already a merged group title, so this leaves existing
        # Mass/Component/Atom child headings and values untouched.
        sheet.cell(1, 2).value = "PINN pass rate (|r| <= 5%)"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", default="outputs/0819final/Final2_checked.xlsx")
    parser.add_argument(
        "--physics-source",
        default="outputs/0819final/three_server_aggregate/physics_conservation_all_runs.csv",
    )
    args = parser.parse_args()
    workbook, source = Path(args.workbook), Path(args.physics_source)
    if not workbook.is_file() or not source.is_file():
        raise FileNotFoundError(f"workbook={workbook}; physics_source={source}")
    summary = _build_summary(source)
    book = openpyxl.load_workbook(workbook)
    _clarify_legacy_pass_rate_headers(book)
    _write_summary_sheet(book, summary)
    temporary = workbook.with_name(f".{workbook.name}.physics.tmp")
    book.save(temporary)
    os.replace(temporary, workbook)
    print(f"updated={workbook}; summary_rows={len(summary)}; source={source}")


if __name__ == "__main__":
    main()
