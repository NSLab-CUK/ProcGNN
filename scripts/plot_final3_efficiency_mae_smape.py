#!/usr/bin/env python3
"""Plot multi-process error versus joint-training time with fold SD bars."""
from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator, ScalarFormatter
import numpy as np
import pandas as pd
from openpyxl import load_workbook


MODEL_ORDER = [
    "Reusable-Distillation-ANN",
    "GCN",
    "GIN",
    "GAT",
    "Graphormer",
    "SAT",
    "GraphToSFILES",
    "Proposed",
]

MODEL_STYLE = {
    "Proposed": {"color": "#000000", "marker": "*", "size": 12.0},
    "GCN": {"color": "#0072B2", "marker": "o", "size": 7.0},
    "GIN": {"color": "#E69F00", "marker": "s", "size": 6.8},
    "GAT": {"color": "#009E73", "marker": "^", "size": 7.2},
    "Graphormer": {"color": "#CC79A7", "marker": "D", "size": 6.6},
    "SAT": {"color": "#D55E00", "marker": "v", "size": 7.2},
    "GraphToSFILES": {"color": "#56B4E9", "marker": "P", "size": 7.0},
    "Reusable-Distillation-ANN": {"color": "#6B7280", "marker": "X", "size": 7.2},
}

FIGSIZE = (8.4, 5.2)

PAIR = re.compile(
    r"^\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)"
    r"(?:\s*(?:±|\+/-)\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?))?"
)


def _pair(value: object) -> tuple[float, float]:
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value), math.nan
    if not isinstance(value, str):
        return math.nan, math.nan
    match = PAIR.match(value)
    if not match:
        return math.nan, math.nan
    return float(match.group(1)), float(match.group(2)) if match.group(2) else math.nan


def _smape_table(workbook: Path) -> pd.DataFrame:
    book = load_workbook(workbook, data_only=True, read_only=False)
    sheet = book["multiprocess_10D"]
    average_smape_column = None
    for column in range(1, sheet.max_column + 1):
        if sheet.cell(1, column).value == "Average" and sheet.cell(2, column + 2).value == "sMAPE (%)":
            average_smape_column = column + 2
            break
    if average_smape_column is None:
        raise RuntimeError("Average sMAPE column not found in multiprocess_10D")

    rows: list[dict[str, object]] = []
    for row in range(3, sheet.max_row + 1):
        model = sheet.cell(row, 1).value
        if model not in MODEL_ORDER:
            continue
        mean, std = _pair(sheet.cell(row, average_smape_column).value)
        rows.append({"Model": model, "sMAPE (%)": mean, "sMAPE SD (%)": std})
    return pd.DataFrame(rows)


def _load_data(workbook: Path, efficiency_csv: Path) -> pd.DataFrame:
    efficiency = pd.read_csv(efficiency_csv)
    efficiency_columns = [
        "Model",
        "Training time (s)",
        "Training time SD (s)",
        "Target-edge mean MAE",
        "Target-edge mean MAE SD",
    ]
    missing = set(efficiency_columns).difference(efficiency.columns)
    if missing:
        raise ValueError(f"efficiency CSV missing columns: {sorted(missing)}")
    frame = efficiency[efficiency_columns].merge(_smape_table(workbook), on="Model", how="inner")
    frame = frame[frame["Model"].isin(MODEL_ORDER)].copy()
    for column in frame.columns[1:]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame = frame.dropna(subset=["Training time (s)", "Target-edge mean MAE", "sMAPE (%)"])
    frame["_order"] = frame["Model"].map({name: idx for idx, name in enumerate(MODEL_ORDER)})
    frame = frame.sort_values("_order").drop(columns="_order").reset_index(drop=True)
    missing_models = sorted(set(MODEL_ORDER).difference(frame["Model"]))
    if missing_models:
        raise RuntimeError(f"missing models after joining performance and efficiency data: {missing_models}")
    return frame


def _save(fig: plt.Figure, stem: Path, dpi: int) -> None:
    # Fixed canvas: all paper figures retain the same aspect ratio after save.
    fig.savefig(stem.with_suffix(".png"), dpi=dpi, facecolor="white")
    fig.savefig(stem.with_suffix(".pdf"), facecolor="white")
    fig.savefig(stem.with_suffix(".svg"), facecolor="white")
    plt.close(fig)


def _plot(frame: pd.DataFrame, mean_column: str, sd_column: str, ylabel: str, stem: Path, dpi: int) -> None:
    fig, axis = plt.subplots(figsize=FIGSIZE)
    axis.set_xscale("log")
    axis.set_axisbelow(True)
    axis.grid(True, which="major", color="#D8DEE6", linewidth=0.8, alpha=0.85)
    axis.grid(True, which="minor", axis="x", color="#E9EDF2", linewidth=0.55, alpha=0.65)
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.spines["left"].set_color("#8A949F")
    axis.spines["bottom"].set_color("#8A949F")
    axis.tick_params(axis="both", labelsize=9, colors="#303640")

    y_lower: list[float] = []
    y_upper: list[float] = []
    for _, row in frame.iterrows():
        model = str(row["Model"])
        x = float(row["Training time (s)"])
        xerr = float(row["Training time SD (s)"])
        y = float(row[mean_column])
        yerr = float(row[sd_column])
        if not math.isfinite(xerr):
            xerr = 0.0
        if not math.isfinite(yerr):
            yerr = 0.0
        style = MODEL_STYLE[model]
        axis.errorbar(
            x,
            y,
            xerr=xerr,
            yerr=yerr,
            fmt=style["marker"],
            color=style["color"],
            ecolor=style["color"],
            markersize=style["size"],
            markeredgecolor="white" if model != "Proposed" else "#000000",
            markeredgewidth=0.8,
            elinewidth=1.45,
            capsize=4.0,
            capthick=1.35,
            linestyle="none",
            label=model,
            zorder=5 if model == "Proposed" else 3,
        )
        y_lower.append(y - yerr)
        y_upper.append(y + yerr)

    axis.set_xlabel("Training time (s)")
    axis.set_ylabel(ylabel)
    axis.yaxis.set_major_locator(MaxNLocator(nbins=6))
    if mean_column == "Target-edge mean MAE":
        formatter = ScalarFormatter(useMathText=True)
        formatter.set_scientific(True)
        formatter.set_powerlimits((3, 3))
        formatter.set_useOffset(False)
        axis.yaxis.set_major_formatter(formatter)
    else:
        axis.yaxis.set_major_formatter(plt.FormatStrFormatter("%.0f"))

    low = max(0.0, min(y_lower))
    high = max(y_upper)
    pad = 0.08 * max(high - low, 1.0)
    axis.set_ylim(max(0.0, low - pad), high + pad)
    axis.legend(
        loc="center left",
        bbox_to_anchor=(1.02, 0.5),
        frameon=False,
        fontsize=8.5,
        handletextpad=0.6,
        borderaxespad=0,
    )
    fig.subplots_adjust(left=0.12, right=0.72, bottom=0.15, top=0.96)
    _save(fig, stem, dpi)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", default="outputs/0819final/Paper_Tables_Final3.xlsx")
    parser.add_argument(
        "--efficiency-csv",
        default="outputs/0819final/Final3_computational_efficiency_mae_nmae/tables/computational_efficiency_multi_process.csv",
    )
    parser.add_argument(
        "--output-dir",
        default="outputs/0819final/Paper_Tables_Final3_computational_efficiency",
    )
    parser.add_argument("--dpi", type=int, default=600)
    args = parser.parse_args()

    workbook = Path(args.workbook)
    efficiency_csv = Path(args.efficiency_csv)
    output = Path(args.output_dir)
    if not workbook.is_file():
        raise FileNotFoundError(workbook)
    if not efficiency_csv.is_file():
        raise FileNotFoundError(efficiency_csv)
    output.mkdir(parents=True, exist_ok=True)

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 10,
            "axes.labelsize": 10,
            "axes.linewidth": 0.8,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )

    frame = _load_data(workbook, efficiency_csv)
    frame.round(4).to_csv(output / "computational_efficiency_mae_smape.csv", index=False, encoding="utf-8-sig")
    _plot(
        frame,
        "Target-edge mean MAE",
        "Target-edge mean MAE SD",
        "MAE",
        output / "mae_vs_training_time",
        args.dpi,
    )
    _plot(
        frame,
        "sMAPE (%)",
        "sMAPE SD (%)",
        "sMAPE (%)",
        output / "smape_vs_training_time",
        args.dpi,
    )
    print(f"models={len(frame)}")
    print(f"output={output.resolve()}")
    print(f"figures=2; formats=png,pdf,svg; dpi={args.dpi}")


if __name__ == "__main__":
    main()
