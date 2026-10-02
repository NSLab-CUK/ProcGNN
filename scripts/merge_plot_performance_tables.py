"""Merge sensitivity/transfer plot data and target-edge sMAPE into a paper workbook.

The MAE/NMAE columns are copied from the Final3 target-performance plot tables.
Whenever saved predictions or post-hoc outputs exist, sMAPE is measured directly.
Only the layer-4 main/default value is taken from the explicitly documented
layer-5-to-layer-4 estimate already stored in Paper_Tables_Final.xlsx.
"""

from __future__ import annotations

import argparse
import math
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from copy import copy
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


PROPERTIES = [
    "Temp", "Pres", "Frac_H2O", "Frac_H2", "Frac_CH4", "Frac_CO2",
    "Frac_CO", "Frac_O2", "Frac_N2", "Mass_Flow",
]
DISPLAY_PROPERTY = {
    "Temp": "Temp", "Pres": "Pres", "Frac_H2O": "H₂O", "Frac_H2": "H₂",
    "Frac_CH4": "CH₄", "Frac_CO2": "CO₂", "Frac_CO": "CO",
    "Frac_O2": "O₂", "Frac_N2": "N₂", "Mass_Flow": "Mass Flow",
}
DISPLAY_TO_RAW = {value: key for key, value in DISPLAY_PROPERTY.items()}
TRANSFER_MODELS = ["GCN", "GIN", "GAT", "Graphormer", "SAT", "GraphToSFILES", "Proposed"]
RATIOS = [0, 2, 4, 6, 8, 10, 20, 30, 40, 50, 60, 70, 80, 90]


@dataclass(frozen=True)
class PredictionJob:
    model: str
    ratio: int
    process: int
    fold: int
    path: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--plot-root", required=True)
    parser.add_argument("--server-root", action="append", required=True)
    parser.add_argument("--workers", type=int, default=10)
    parser.add_argument("--cache-csv", required=True)
    return parser.parse_args()


def mean_std(values: list[float]) -> tuple[float, float]:
    array = np.asarray(values, dtype=float)
    return float(array.mean()), float(array.std(ddof=1)) if len(array) > 1 else 0.0


def parse_process_fold(path: Path) -> tuple[int, int]:
    text = str(path).replace("\\", "/")
    process = int(re.search(r"heldout_P(\d+)", text).group(1))
    fold = int(re.search(r"fold_(\d+)", text).group(1))
    return process, fold


def discover_baseline_jobs(server_roots: list[Path]) -> list[PredictionJob]:
    jobs: dict[tuple[str, int, int, int], PredictionJob] = {}
    for root in server_roots:
        final = root / "outputs" / "0819final"
        transfer = final / "baselines_transfer_0909"
        zero = final / "baselines_zero_shot_0910"
        for model in TRANSFER_MODELS:
            if model == "Proposed":
                continue
            for path in (transfer / model).glob("heldout_P*/fold_*/ratio_*/test_predictions.csv"):
                process, fold = parse_process_fold(path)
                ratio = int(path.parent.name.split("_")[-1])
                jobs[(model, ratio, process, fold)] = PredictionJob(model, ratio, process, fold, path)
            for path in (zero / model).glob("heldout_P*/fold_*/test_predictions.csv"):
                process, fold = parse_process_fold(path)
                jobs[(model, 0, process, fold)] = PredictionJob(model, 0, process, fold, path)
    expected = 6 * len(RATIOS) * 10 * 5
    if len(jobs) != expected:
        raise RuntimeError(f"Expected {expected} unique baseline transfer/zero-shot jobs, found {len(jobs)}")
    return sorted(jobs.values(), key=lambda x: (x.model, x.ratio, x.process, x.fold))


def read_prediction(job: PredictionJob) -> list[dict[str, object]]:
    frame = pd.read_csv(job.path, usecols=["property_name", "true_value", "predicted_value"])
    frame["true_value"] = pd.to_numeric(frame["true_value"], errors="coerce")
    frame["predicted_value"] = pd.to_numeric(frame["predicted_value"], errors="coerce")
    true = frame["true_value"].to_numpy(float)
    pred = frame["predicted_value"].to_numpy(float)
    valid = np.isfinite(true) & np.isfinite(pred)
    frame = frame.loc[valid, ["property_name"]].copy()
    frame["term"] = 100.0 * 2.0 * np.abs(pred[valid] - true[valid]) / (
        np.abs(pred[valid]) + np.abs(true[valid]) + 1.0e-8
    )
    grouped = frame.groupby("property_name", sort=False)["term"].agg(["mean", "count"])
    rows = []
    for prop in PROPERTIES:
        if prop not in grouped.index:
            raise RuntimeError(f"Missing {prop} in {job.path}")
        rows.append({
            "model": job.model, "transfer_percent": job.ratio, "process": job.process,
            "fold": job.fold, "property_name": prop,
            "sMAPE_pct": float(grouped.loc[prop, "mean"]),
            "n": int(grouped.loc[prop, "count"]), "source_path": str(job.path),
            "source_kind": "baseline_prediction_csv",
        })
    return rows


def load_or_compute_baseline(jobs: list[PredictionJob], cache: Path, workers: int) -> pd.DataFrame:
    if cache.exists():
        frame = pd.read_csv(cache)
        keys = frame[["model", "transfer_percent", "process", "fold"]].drop_duplicates()
        if len(keys) == len(jobs) and len(frame) == len(jobs) * len(PROPERTIES):
            print(f"[cache] using {cache} ({len(frame)} rows)", flush=True)
            return frame
    all_rows: list[dict[str, object]] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        for index, rows in enumerate(pool.map(read_prediction, jobs), start=1):
            all_rows.extend(rows)
            if index % 100 == 0 or index == len(jobs):
                print(f"[baseline transfer sMAPE] {index}/{len(jobs)}", flush=True)
    frame = pd.DataFrame(all_rows)
    cache.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(cache, index=False, encoding="utf-8-sig")
    return frame


def load_posthoc_file(path: Path, model: str, ratio: int, process: int, fold: int) -> list[dict[str, object]]:
    frame = pd.read_csv(path, usecols=["property_name", "sMAPE_pct"])
    return [
        {
            "model": model, "transfer_percent": ratio, "process": process, "fold": fold,
            "property_name": str(row.property_name), "sMAPE_pct": float(row.sMAPE_pct),
            "n": math.nan, "source_path": str(path), "source_kind": "proposed_posthoc",
        }
        for row in frame.itertuples(index=False)
    ]


def proposed_transfer(server_roots: list[Path]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    seen: set[tuple[int, int, int]] = set()
    for root in server_roots:
        base = root / "outputs" / "0819final" / "posthoc_target_smape"
        for path in (base / "data_efficiency").glob("heldout_P*/fold_*/transfer/ratio_*/target_property_smape.csv"):
            process, fold = parse_process_fold(path)
            ratio = int(path.parent.name.split("_")[-1])
            key = (ratio, process, fold)
            if key not in seen:
                seen.add(key)
                rows.extend(load_posthoc_file(path, "Proposed", ratio, process, fold))
        for path in (base / "proposed_unseen").glob("heldout_P*/fold_*/zero_shot/target_property_smape.csv"):
            process, fold = parse_process_fold(path)
            key = (0, process, fold)
            if key not in seen:
                seen.add(key)
                rows.extend(load_posthoc_file(path, "Proposed", 0, process, fold))
    expected = len(RATIOS) * 10 * 5
    if len(seen) != expected:
        raise RuntimeError(f"Expected {expected} Proposed transfer/zero-shot jobs, found {len(seen)}")
    return pd.DataFrame(rows)


def aggregate_transfer(detail: pd.DataFrame) -> pd.DataFrame:
    groups = []
    for (model, ratio, prop), part in detail.groupby(["model", "transfer_percent", "property_name"]):
        mean, std = mean_std(part["sMAPE_pct"].tolist())
        groups.append({"model": model, "transfer_percent": int(ratio), "property_name": prop,
                       "sMAPE_mean_pct": mean, "sMAPE_std_pct": std,
                       "sMAPE_status": "measured"})
    # Average follows the plot convention: arithmetic mean of the ten property metrics per run.
    piv = detail.pivot_table(index=["model", "transfer_percent", "process", "fold"],
                             columns="property_name", values="sMAPE_pct", aggfunc="first")
    piv["Average"] = piv[PROPERTIES].mean(axis=1)
    for (model, ratio), part in piv.groupby(level=[0, 1]):
        mean, std = mean_std(part["Average"].tolist())
        groups.append({"model": model, "transfer_percent": int(ratio), "property_name": "Average",
                       "sMAPE_mean_pct": mean, "sMAPE_std_pct": std,
                       "sMAPE_status": "measured"})
    return pd.DataFrame(groups)


def posthoc_sensitivity(server_roots: list[Path]) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    seen: set[str] = set()
    for root in server_roots:
        base = root / "outputs" / "0819final" / "posthoc_target_smape"
        paths = list((base / "sensitivity_10d_clean").glob("**/target_property_smape.csv"))
        paths += list((base / "proposed_joint_10d_clean").glob("All/fold_*/target_property_smape.csv"))
        for path in paths:
            normalized = str(path).replace("\\", "/")
            fold = int(re.search(r"fold_(\d+)", normalized).group(1))
            if "/depth/depth_" in normalized:
                setting = "Depth " + re.search(r"/depth/depth_(\d+)/", normalized).group(1)
            elif "/pin/node_" in normalized:
                token = re.search(r"/pin/node_([^/]+)/", normalized).group(1)
                family, level = token.rsplit("_", 1)
                setting = family.title() + " " + level.title()
            elif "/proposed_joint_10d_clean/" in normalized:
                setting = "Depth 5"
            else:
                continue
            key = f"{setting}|{fold}"
            if key in seen:
                continue
            seen.add(key)
            frame = pd.read_csv(path, usecols=["property_name", "sMAPE_pct"])
            for row in frame.itertuples(index=False):
                records.append({"setting_key": setting, "fold": fold,
                                "property_name": str(row.property_name),
                                "sMAPE_pct": float(row.sMAPE_pct), "source_path": str(path)})
    expected = 13 * 5  # six settings + depth 5 + six PIN low/high settings
    if len(seen) != expected:
        raise RuntimeError(f"Expected {expected} sensitivity jobs, found {len(seen)}: {sorted(seen)}")
    return pd.DataFrame(records)


def paper_l4_default(workbook: Path) -> tuple[float, float]:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["multiprocess_10D"]
    smape_cols = [c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "sMAPE (%)"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Proposed")
    value = str(ws.cell(row, smape_cols[-1]).value)
    nums = re.findall(r"[-+]?\d+(?:\.\d+)?", value)
    if len(nums) < 2:
        raise RuntimeError(f"Cannot parse Proposed Average sMAPE from {value!r}")
    return float(nums[0]), float(nums[1])


def build_sensitivity_table(source_csv: Path, posthoc: pd.DataFrame, workbook: Path) -> pd.DataFrame:
    source = pd.read_csv(source_csv)
    wide = source.pivot_table(index=["sensitivity", "setting", "setting_order"],
                              columns="metric", values="mean", aggfunc="first").reset_index()
    weights = {
        ("Mass", "Low"): 0.5, ("Mass", "Default"): 1.0, ("Mass", "High"): 2.0,
        ("Component", "Low"): 7.5e-8, ("Component", "Default"): 1.5e-7,
        ("Component", "High"): 3.0e-7, ("Atom", "Low"): 0.1,
        ("Atom", "Default"): 0.2, ("Atom", "High"): 0.4,
    }
    l4_mean, l4_std = paper_l4_default(workbook)
    stats: dict[str, tuple[float, float]] = {}
    for setting, part in posthoc.groupby("setting_key"):
        piv = part.pivot(index="fold", columns="property_name", values="sMAPE_pct")
        macro = piv[PROPERTIES].mean(axis=1)
        stats[setting] = mean_std(macro.tolist())

    out = []
    for row in wide.itertuples(index=False):
        sensitivity, setting = str(row.sensitivity), str(row.setting)
        key = f"Depth {setting}" if sensitivity == "Depth" else f"{sensitivity} {setting}"
        if setting == "Default" and sensitivity in {"Mass", "Component", "Atom"}:
            smape_mean, smape_std = l4_mean, l4_std
            status = "estimated layer-4 from layer-5 correction (dagger in main table)"
        else:
            smape_mean, smape_std = stats[key]
            status = "measured layer-5 posthoc" if sensitivity != "Depth" else "measured depth-specific posthoc"
        out.append({
            "sensitivity": sensitivity, "setting": setting,
            "weight_or_depth": weights.get((sensitivity, setting), int(setting) if sensitivity == "Depth" else math.nan),
            "setting_order": int(row.setting_order), "MAE_mean": float(row.MAE),
            "NMAE_mean": float(row.NMAE), "sMAPE_mean_pct": smape_mean,
            "sMAPE_std_pct": smape_std, "sMAPE_status": status,
        })
    order = {"Depth": 0, "Mass": 1, "Component": 2, "Atom": 3}
    result = pd.DataFrame(out)
    result["_order"] = result["sensitivity"].map(order)
    return result.sort_values(["_order", "setting_order"]).drop(columns="_order").reset_index(drop=True)


def build_transfer_table(property_csv: Path, average_csv: Path, smape: pd.DataFrame) -> pd.DataFrame:
    prop = pd.read_csv(property_csv)
    mean = prop.pivot_table(index=["model", "transfer_percent", "property"], columns="metric",
                            values="mean", aggfunc="first").reset_index()
    std = prop.pivot_table(index=["model", "transfer_percent", "property"], columns="metric",
                           values="std", aggfunc="first").reset_index()
    mean = mean.rename(columns={"MAE": "MAE_mean", "NMAE": "NMAE_mean"})
    std = std.rename(columns={"MAE": "MAE_std", "NMAE": "NMAE_std"})
    table = mean.merge(std, on=["model", "transfer_percent", "property"], how="outer")
    avg = pd.read_csv(average_csv).pivot_table(index=["model", "transfer_percent"], columns="metric",
                                               values="mean", aggfunc="first").reset_index()
    avg = avg.rename(columns={"MAE": "MAE_mean", "NMAE": "NMAE_mean"})
    avg["property"] = "Average"
    avg["MAE_std"] = math.nan
    avg["NMAE_std"] = math.nan
    table = pd.concat([table, avg[table.columns]], ignore_index=True)
    smape = smape.copy()
    smape["property"] = smape["property_name"].map(DISPLAY_PROPERTY).fillna(smape["property_name"])
    smape = smape.drop(columns="property_name")
    table = table.merge(smape, on=["model", "transfer_percent", "property"], how="left", validate="one_to_one")
    if table["sMAPE_mean_pct"].isna().any():
        missing = table.loc[table["sMAPE_mean_pct"].isna(), ["model", "transfer_percent", "property"]]
        raise RuntimeError(f"Missing transfer sMAPE rows:\n{missing.to_string(index=False)}")
    model_order = {name: index for index, name in enumerate(TRANSFER_MODELS)}
    property_order = {DISPLAY_PROPERTY[p]: i for i, p in enumerate(PROPERTIES)} | {"Average": len(PROPERTIES)}
    table["_m"] = table.model.map(model_order)
    table["_p"] = table.property.map(property_order)
    cols = ["model", "transfer_percent", "property", "MAE_mean", "MAE_std", "NMAE_mean", "NMAE_std",
            "sMAPE_mean_pct", "sMAPE_std_pct", "sMAPE_status"]
    return table.sort_values(["_m", "transfer_percent", "_p"])[cols].reset_index(drop=True)


def style_sheet(ws, *, freeze: str = "A2") -> None:
    dark = PatternFill("solid", fgColor="1F4E78")
    light = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7B7B7")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = dark
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 100) + 1)]
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max(max(map(len, values)) + 2, 11), 54)
    ws.row_dimensions[1].height = 32


def write_dataframe(wb, name: str, frame: pd.DataFrame) -> None:
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    style_sheet(ws)


def write_notes(wb, baseline_jobs: int, proposed_jobs: int) -> None:
    name = "Plot_Data_Notes"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    rows = [
        ("Item", "Detail"),
        ("Scope", "Target edges only; All-edge results are excluded."),
        ("sMAPE formula", "100 * mean(2*abs(prediction-truth)/(abs(prediction)+abs(truth)+1e-8)); lower is better."),
        ("Sensitivity table", "MAE/NMAE copied from Final3_target_performance plot data; sMAPE added from saved posthoc outputs."),
        ("Transfer table", "MAE/NMAE copied from Final3 target plot data; sMAPE measured from saved predictions/posthoc outputs for 0%-90%."),
        ("Layer convention", "Depth rows use their named depth. PIN Low/High and transfer sMAPE are measured layer-5 outputs."),
        ("Estimated values", "PIN Default uses the main layer-4 sMAPE estimate already documented in sMAPE_L4_Estimation."),
        ("Aggregation", "Property values: mean and sample SD across process-fold runs. Average: arithmetic mean of 10 property metrics per run, then mean/SD."),
        ("Baseline jobs", baseline_jobs),
        ("Proposed transfer jobs", proposed_jobs),
        ("Rounding", "Numeric cells retain full precision and display four decimal places."),
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    args = parse_args()
    workbook = Path(args.workbook).resolve()
    plot_root = Path(args.plot_root).resolve()
    roots = [Path(item).resolve() for item in args.server_root]
    cache = Path(args.cache_csv).resolve()

    jobs = discover_baseline_jobs(roots)
    baseline = load_or_compute_baseline(jobs, cache, args.workers)
    proposed = proposed_transfer(roots)
    detail = pd.concat([baseline, proposed], ignore_index=True)
    transfer_smape = aggregate_transfer(detail)
    transfer = build_transfer_table(
        plot_root / "tables" / "transfer_target_property_performance.csv",
        plot_root / "tables" / "transfer_target_average_performance.csv",
        transfer_smape,
    )
    sensitivity_posthoc = posthoc_sensitivity(roots)
    sensitivity = build_sensitivity_table(
        plot_root / "tables" / "sensitivity_target_mean_performance.csv",
        sensitivity_posthoc,
        workbook,
    )

    wb = openpyxl.load_workbook(workbook)
    write_dataframe(wb, "Plot_Sensitivity", sensitivity)
    write_dataframe(wb, "Plot_Transfer", transfer)
    write_notes(wb, len(jobs), len(proposed) // len(PROPERTIES))
    temp = workbook.with_name(workbook.stem + ".plot_merge_tmp.xlsx")
    wb.save(temp)
    temp.replace(workbook)
    print(f"[done] workbook={workbook}", flush=True)
    print(f"[done] sensitivity_rows={len(sensitivity)} transfer_rows={len(transfer)}", flush=True)
    print(f"[done] baseline_jobs={len(jobs)} proposed_jobs={len(proposed)//len(PROPERTIES)}", flush=True)


if __name__ == "__main__":
    main()
