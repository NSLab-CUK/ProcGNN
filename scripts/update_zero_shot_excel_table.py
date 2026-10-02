#!/usr/bin/env python3
"""Insert completed zero-shot baseline rows into the paper Excel workbook."""
from __future__ import annotations

import argparse
import os
from copy import copy
from pathlib import Path

import openpyxl
import pandas as pd

from aggregate_three_server_phase_metrics import PROPERTIES, _build_model_metric_tables, _with_macro_stats


BASELINES = ["GCN", "GIN", "GAT", "Graphormer", "SAT", "GraphToSFILES"]


def _zero_shot_table(aggregate: Path) -> pd.DataFrame:
    prop = pd.read_csv(aggregate / "phase_family_property_summary.csv")
    summary = pd.read_csv(aggregate / "phase_family_summary.csv")
    table = _build_model_metric_tables(
        _with_macro_stats(prop, summary), formatted_mean_std=True,
    )["zero_shot"]
    required = {f"{name}-zero-shot" for name in BASELINES}
    present = set(table["Model/Condition"].astype(str))
    missing = sorted(required - present)
    if missing:
        raise RuntimeError(f"missing completed baseline zero-shot rows: {missing}")
    return table


def _copy_style(source, destination) -> None:
    if source.has_style:
        destination._style = copy(source._style)
    if source.number_format:
        destination.number_format = source.number_format
    if source.alignment:
        destination.alignment = copy(source.alignment)
    if source.protection:
        destination.protection = copy(source.protection)


def _excel_values(table: pd.DataFrame, model: str) -> list[str]:
    display = f"{model}-zero-shot"
    rows = table[table["Model/Condition"].astype(str).eq(display)]
    mae = rows[rows["Metric"].astype(str).eq("MAE")].iloc[0]
    smae = rows[rows["Metric"].astype(str).eq("sMAE")].iloc[0]
    values = [f"{model} (Target edge)"]
    for property_name in PROPERTIES:
        values.extend([str(mae[property_name]), str(smae[property_name])])
    values.extend([str(mae["Mean"]), str(smae["Mean"])])
    return values


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--aggregate-dir", default="outputs/0819final/three_server_aggregate")
    parser.add_argument("--workbook", default="outputs/0819final/ALL_EXPERIMENT_TABLES_0910_Final.xlsx")
    args = parser.parse_args()
    aggregate, workbook = Path(args.aggregate_dir), Path(args.workbook)
    table = _zero_shot_table(aggregate)

    book = openpyxl.load_workbook(workbook)
    if "Zero_Shot" not in book.sheetnames:
        raise KeyError("workbook has no Zero_Shot sheet")
    sheet = book["Zero_Shot"]
    source_row = 3  # Existing Proposed target-edge row defines the target-edge styling.
    if sheet.cell(source_row, 1).value != "Proposed (Target edge)":
        raise RuntimeError("unexpected Zero_Shot layout: Proposed target-edge row is not row 3")

    labels = {f"{model} (Target edge)" for model in BASELINES}
    for row in range(sheet.max_row, 2, -1):
        if str(sheet.cell(row, 1).value) in labels:
            sheet.delete_rows(row, 1)

    insert_at = next(
        (row for row in range(3, sheet.max_row + 1) if sheet.cell(row, 1).value == "Proposed (All edge)"),
        sheet.max_row + 1,
    )
    sheet.insert_rows(insert_at, amount=len(BASELINES))
    for offset, model in enumerate(BASELINES):
        row = insert_at + offset
        for column, value in enumerate(_excel_values(table, model), start=1):
            target = sheet.cell(row, column)
            _copy_style(sheet.cell(source_row, column), target)
            target.value = value
        sheet.row_dimensions[row].height = sheet.row_dimensions[source_row].height

    temporary = workbook.with_name(f".{workbook.name}.zero_shot.tmp")
    book.save(temporary)
    os.replace(temporary, workbook)
    print(f"updated {workbook}: inserted {len(BASELINES)} baseline zero-shot rows")


if __name__ == "__main__":
    main()
