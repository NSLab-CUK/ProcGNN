from __future__ import annotations

import hashlib
import sys
from dataclasses import replace
from pathlib import Path

import pandas as pd
import torch
import torch.nn.functional as F
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from process_graph.data.tabular_dataset import (  # noqa: E402
    ProcessGraphTabularDataset,
    collate_graph_batch,
)
from process_graph.experiment.config_builders import (  # noqa: E402
    build_task_specs,
    model_yaml_to_encoder_config,
)
from process_graph.experiment.edge_step_pi_training import (  # noqa: E402
    build_pi_physical_outputs,
)
from process_graph.experiment.loaders import load_experiment_config  # noqa: E402
from process_graph.experiment.node_balance_pi import (  # noqa: E402
    compute_node_balance_pinn_losses,
)
from process_graph.models.process_surrogate import ProcessSurrogateModel  # noqa: E402


PROCESS1_STREAM_TO_EDGE = {
    "14": ("SP1", "P1"),
    "15": ("SP1", "P2"),
    "16": ("SP1", "P3"),
    "19": ("T1", "MIX2"),
    "22": ("HX3", "MIX2"),
    "24": ("HX4", "MIX2"),
}

UNCHANGED_PROCESS_EDGE_HASHES = {
    2: "ae8d2bb6819f6c3e3d97b59fc6e472c8863519f789e714d862854b8333bd57f1",
    3: "c588a015df568577d9990d008891bc1d829a01b6716cafe43739f21efa87fd45",
    4: "db190e38affa3f0ec34fdcc38d5a0bdb0cf1349335a6f16ad9afdd14b859d968",
    5: "baf1913ea6041f6caee92157a5e09340efa9808de64eae0c8379a4c1cee92565",
    6: "bc068e80d5f39932be4066eea2ee7b3b63ae4a5a583d44e0d4bd39133f8fa74b",
    7: "5c987e2588b83a70bf043d9f8227b3ba76d79dd6a9690630496b177382c22d4f",
    8: "548ae414ea6429c5302cda57b0a61259b176c9cee247a8fb5516dfa1523d548b",
    9: "19b1914f553599b572dc6959ab0e2048f010bbc3647dfab98a06c24cd12ea5a3",
    10: "707d44d6d1d670d7e32210cf9d970874635f7c5c3b866e1f2d65e3a4d97ce00f",
}


def _matrix_edges(path: Path) -> set[tuple[str, str]]:
    workbook = load_workbook(path, data_only=False, read_only=True)
    sheet = workbook.worksheets[0]
    destinations = {
        column: str(sheet.cell(1, column).value).strip()
        for column in range(2, sheet.max_column + 1)
    }
    edges: set[tuple[str, str]] = set()
    for row in range(2, sheet.max_row + 1):
        source = str(sheet.cell(row, 1).value).strip()
        for column, destination in destinations.items():
            value = sheet.cell(row, column).value
            if value is not None and float(value) != 0.0:
                edges.add((source, destination))
    return edges


def _has_path(edges: set[tuple[str, str]], source: str, destination: str) -> bool:
    frontier = [source]
    visited: set[str] = set()
    while frontier:
        node = frontier.pop()
        if node == destination:
            return True
        if node in visited:
            continue
        visited.add(node)
        frontier.extend(dst for src, dst in edges if src == node)
    return False


def test_process1_source_workbooks_and_canonical_topology() -> None:
    workbook_paths = (
        PROJECT_ROOT / "data/adjacency_matrix/Process1_Adjacency_Matrix.xlsx",
        PROJECT_ROOT / "data/process_specs/raw/Process1_Adjacency_Matrix.xlsx",
    )
    expected_sp1_out = {"P1", "P2", "P3"}
    expected_mix2_in = {"T1", "HX3", "HX4"}
    for path in workbook_paths:
        edges = _matrix_edges(path)
        assert {dst for src, dst in edges if src == "SP1"} == expected_sp1_out
        assert {src for src, dst in edges if dst == "MIX2"} == expected_mix2_in
        assert ("SP1", "MIX2") not in edges

    edge_path = PROJECT_ROOT / "data/reference/v3/canonical_edges.csv"
    frame = pd.read_csv(edge_path, dtype=str, keep_default_na=False)
    p1 = frame[frame["process_id"].eq("1")].copy()
    topology = set(zip(p1["src_node"], p1["dst_node"]))

    assert len(p1) == 38
    assert "P01_E032" not in set(p1["canonical_edge_id"])
    assert {dst for src, dst in topology if src == "SP1"} == expected_sp1_out
    assert {src for src, dst in topology if dst == "MIX2"} == expected_mix2_in
    assert ("SP1", "MIX2") not in topology

    for stream_key, expected_nodes in PROCESS1_STREAM_TO_EDGE.items():
        row = p1[p1["main_data_stream_key"].eq(stream_key)]
        assert len(row) == 1
        assert (row.iloc[0]["src_node"], row.iloc[0]["dst_node"]) == expected_nodes

    stable_ids = {
        "14": "P01_E033",
        "15": "P01_E034",
        "16": "P01_E035",
        "19": "P01_E029",
        "22": "P01_E017",
        "24": "P01_E019",
    }
    for stream_key, edge_id in stable_ids.items():
        assert p1.loc[
            p1["main_data_stream_key"].eq(stream_key), "canonical_edge_id"
        ].item() == edge_id

    outputs = p1[p1["dst_node"].eq("V_OUTPUT")]
    assert set(
        zip(outputs["src_node"], outputs["main_data_stream_key"])
    ) == {
        ("HX2", "FUELGAS"),
        ("HX5", "PROD"),
        ("T3", "RESTEAM"),
    }
    assert not p1.duplicated(
        subset=["src_node", "dst_node", "main_data_stream_key"]
    ).any()

    for node in ("P1", "P2", "P3", "MIX2"):
        assert any(src == node or dst == node for src, dst in topology)
        assert _has_path(topology, "V_INPUT", node)
        assert _has_path(topology, node, "V_OUTPUT")


def test_process2_to_process10_canonical_edges_are_unchanged() -> None:
    frame = pd.read_csv(
        PROJECT_ROOT / "data/reference/v3/canonical_edges.csv",
        dtype=str,
        keep_default_na=False,
    )
    for process_id, expected_hash in UNCHANGED_PROCESS_EDGE_HASHES.items():
        rows = frame[frame["process_id"].eq(str(process_id))].sort_values(
            "canonical_edge_id"
        )
        serialized = rows.to_csv(index=False, lineterminator="\n")
        assert hashlib.sha256(serialized.encode()).hexdigest() == expected_hash


def test_process1_dataset_forward_backward_metric_and_pinn_smoke() -> None:
    experiment = load_experiment_config(
        PROJECT_ROOT
        / "configs/experiment/pinn/known_feed_260730/f0_e2_baseline.yaml"
    )
    data_config = replace(
        experiment.data,
        normalize_x_oper=False,
        normalize_targets=False,
        normalize_y_edge=False,
        edge_all_processes=[1],
    )
    dataset = ProcessGraphTabularDataset(
        experiment.data.train_data_path,
        data_config,
        experiment.project_root,
    )
    process_column = data_config.process_id_column
    sample_index = int(
        dataset.frame.index[
            dataset.frame[process_column].astype(str).eq("Process1")
        ][0]
    )
    record = dataset[sample_index]
    graph = record.graph

    assert len(graph.canonical_edge_ids) == 38
    assert sum(graph.edge_is_supervised) == 38.0
    assert sum(graph.edge_feature_mask) == 38.0
    assert sum(graph.y_edge_mask) == 38.0
    assert sum(graph.edge_pinn_mask) == 38.0
    assert {
        edge_id
        for edge_id, target in zip(graph.canonical_edge_ids, graph.edge_is_target)
        if target > 0.5
    } == {"P01_E015", "P01_E021"}

    def incident(node: str, incoming: bool) -> set[str]:
        node_index = graph.node_names.index(node)
        side = graph.edge_index[1 if incoming else 0]
        return {
            edge_id
            for edge_id, endpoint, pinn_mask in zip(
                graph.canonical_edge_ids, side, graph.edge_pinn_mask
            )
            if endpoint == node_index and pinn_mask > 0.5
        }

    assert incident("SP1", incoming=False) == {
        "P01_E033",
        "P01_E034",
        "P01_E035",
    }
    assert incident("MIX2", incoming=True) == {
        "P01_E017",
        "P01_E019",
        "P01_E029",
    }

    batch = collate_graph_batch([record])
    model = ProcessSurrogateModel(
        model_yaml_to_encoder_config(experiment.model, data_config),
        task_specs=build_task_specs(experiment.model, data_config),
    )
    outputs = model(batch.model_kwargs, task_inputs=batch.task_inputs)
    predictions = outputs["main_stream_pred"]
    assert predictions.shape == (38, 11)
    assert torch.isfinite(predictions).all()

    prediction_columns = [
        "Temp",
        "Pres",
        *[f"Frac_{name}" for name in experiment.model.species_order],
        "Mass_Flow",
        "Vol_Flow",
    ]
    target_columns = list(batch.edge_target_columns)
    target_indices = [target_columns.index(name) for name in prediction_columns]
    targets = batch.model_kwargs["y_edge_true"][:, target_indices]
    valid_rows = batch.model_kwargs["y_edge_mask"].bool()
    edge_loss = F.smooth_l1_loss(predictions[valid_rows], targets[valid_rows])

    physical_outputs = build_pi_physical_outputs(
        outputs,
        train_cfg=experiment.train,
        data_cfg=data_config,
        normalizer=None,
    )
    fraction_indices = [
        target_columns.index(f"Frac_{name}")
        for name in experiment.model.species_order
    ]
    physical_targets = {
        "mass_flow": batch.model_kwargs["y_edge_true"][
            :, [target_columns.index("Mass_Flow")]
        ],
        "frac": batch.model_kwargs["y_edge_true"][:, fraction_indices],
    }
    pinn = compute_node_balance_pinn_losses(
        physical_outputs=physical_outputs,
        physical_targets=physical_targets,
        edge_index=batch.model_kwargs["edge_index"],
        edge_batch_or_graph_id=batch.model_kwargs["edge_batch"],
        node_batch_or_graph_id=batch.model_kwargs["batch"],
        node_unit_type=batch.model_kwargs["x_unit"],
        species_order=experiment.model.species_order,
        train_cfg=experiment.train,
        node_q=batch.model_kwargs["node_q"],
        node_w=batch.model_kwargs["node_w"],
        node_qw_valid_mask=batch.model_kwargs["node_qw_valid_mask"],
        node_balance_exclude_mass=batch.model_kwargs["node_balance_exclude_mass"],
        node_balance_exclude_component=batch.model_kwargs[
            "node_balance_exclude_component"
        ],
        node_balance_exclude_atom=batch.model_kwargs["node_balance_exclude_atom"],
        node_balance_exclude_energy=batch.model_kwargs[
            "node_balance_exclude_energy"
        ],
        edge_pinn_mask=batch.model_kwargs["edge_pinn_mask"],
    )
    assert torch.isfinite(pinn["weighted_node_total"])
    total_loss = edge_loss + pinn["weighted_node_total"]
    model.zero_grad(set_to_none=True)
    total_loss.backward()
    assert any(
        parameter.grad is not None
        for parameter in model.parameters()
        if parameter.requires_grad
    )

    metric_absolute_error = torch.abs(
        predictions[valid_rows].detach() - targets[valid_rows]
    ).mean()
    assert torch.isfinite(metric_absolute_error)
