"""Build a portable Aspen handoff for the expanded 135-candidate GA queue."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import shutil
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--study-root", required=True)
    parser.add_argument("--candidate-dir", required=True)
    parser.add_argument("--validation-workbook", required=True)
    parser.add_argument("--output-dir", required=True)
    return parser.parse_args()


def resolve_path(value: str) -> Path:
    path = Path(value)
    return path.resolve() if path.is_absolute() else (PROJECT_ROOT / path).resolve()


def as_float(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def process_number(value: object) -> int:
    return int(str(value).strip().upper().removeprefix("P"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def unit_hint(name: str) -> str:
    upper = name.upper()
    if upper.startswith("P_") or upper.startswith("PRES"):
        return "bar (confirm in Aspen case)"
    if upper.startswith("T_") or upper.startswith("DT_") or "TEMP" in upper:
        return "degC or delta-degC (confirm in Aspen case)"
    if "FRAC" in upper:
        return "fraction"
    if "FLOW" in upper:
        return "native Process_Main/Aspen flow unit (confirm)"
    return "confirm in original Aspen case"


def relative_error(prediction: object, reference: object) -> float | None:
    prediction_f = as_float(prediction)
    reference_f = as_float(reference)
    if prediction_f is None or reference_f is None:
        return None
    if abs(prediction_f) < 1e-8 and abs(reference_f) < 1e-8:
        return 0.0
    return abs(prediction_f - reference_f) / max(abs(reference_f), 1e-12)


def load_payloads(study_root: Path, candidates: pd.DataFrame) -> dict[str, dict[str, Any]]:
    payloads: dict[str, dict[str, Any]] = {}
    for run_id in sorted(candidates["run_id"].astype(str).unique()):
        path = study_root / run_id / "ga_summary.json"
        if not path.is_file():
            raise FileNotFoundError(path)
        payload = json.loads(path.read_text(encoding="utf-8"))
        if payload.get("input_policy") != "masked_proxy":
            raise ValueError(f"Unsafe input policy in {path}")
        if int(payload.get("arguments", {}).get("template_id", -1)) != 5000:
            raise ValueError(f"Expected template ID=5000 in {path}")
        payloads[run_id] = payload
    return payloads


def process_payload(payload: dict[str, Any], process_id: str) -> dict[str, Any]:
    matches = [
        item for item in payload.get("processes", [])
        if str(item.get("process_id")) == process_id
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one summary for {process_id}, found {len(matches)}")
    return matches[0]


def safe_columns(payload: dict[str, Any], process_id: str) -> list[str]:
    matches = [
        item for item in payload.get("input_audits", [])
        if str(item.get("process_id")) == process_id
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one input audit for {process_id}, found {len(matches)}")
    return [str(value) for value in matches[0].get("safe_condition_columns", [])]


def load_previous_validation(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["결과요약"]
    records: list[dict[str, Any]] = []
    for row_index in range(9, 19):
        process_id = sheet.cell(row_index, 1).value
        if not process_id:
            continue
        gnn_h2, aspen_h2 = sheet.cell(row_index, 8).value, sheet.cell(row_index, 9).value
        gnn_lcoh, aspen_lcoh = sheet.cell(row_index, 12).value, sheet.cell(row_index, 13).value
        gnn_co2, aspen_co2 = sheet.cell(row_index, 16).value, sheet.cell(row_index, 17).value
        records.append(
            {
                "process_id": str(process_id),
                "candidate_available": sheet.cell(row_index, 2).value,
                "aspen_status": sheet.cell(row_index, 3).value,
                "uosstat": sheet.cell(row_index, 4).value,
                "block_error": sheet.cell(row_index, 5).value,
                "input_change_count": sheet.cell(row_index, 6).value,
                "validation_note": sheet.cell(row_index, 7).value,
                "gnn_h2_production_kg_h": gnn_h2,
                "aspen_h2_production_kg_h": aspen_h2,
                "h2_relative_error": relative_error(gnn_h2, aspen_h2),
                "gnn_lcoh_usd_per_kg": gnn_lcoh,
                "aspen_lcoh_usd_per_kg": aspen_lcoh,
                "lcoh_relative_error": relative_error(gnn_lcoh, aspen_lcoh),
                "gnn_specific_co2_kg_per_kg_h2": gnn_co2,
                "aspen_specific_co2_kg_per_kg_h2": aspen_co2,
                "co2_relative_error": relative_error(gnn_co2, aspen_co2),
                "correction_reference_row_not_baseline": sheet.cell(row_index, 20).value,
                "seed": sheet.cell(row_index, 22).value,
                "run_id": sheet.cell(row_index, 23).value,
            }
        )
    return pd.DataFrame(records)


def build_tables(
    study_root: Path,
    candidates: pd.DataFrame,
    payloads: dict[str, dict[str, Any]],
    previous: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    process_ids = sorted(candidates["process_id"].astype(str).unique(), key=process_number)
    previous_by_process = {
        str(row["process_id"]): row for row in previous.to_dict("records")
    }
    run_long = pd.read_csv(study_root / "paper_summary" / "paper_runs_long.csv")
    run_long = run_long.loc[
        (run_long["scenario"] == "LCOH_capacity") & (run_long["search_method"] == "ga")
    ]

    template_cache: dict[str, tuple[pd.Series, list[str], str]] = {}
    baseline_condition_records: list[dict[str, Any]] = []
    baseline_long_records: list[dict[str, Any]] = []
    for process_id in process_ids:
        first = candidates.loc[candidates["process_id"] == process_id].iloc[0]
        payload = payloads[str(first["run_id"])]
        columns = safe_columns(payload, process_id)
        source = PROJECT_ROOT / "data" / "main_data" / f"{process_number(process_id)}.Process_Main.csv"
        frame = pd.read_csv(source)
        selected = frame.loc[pd.to_numeric(frame["ID"], errors="coerce").eq(5000)]
        if len(selected) != 1:
            raise ValueError(f"{source}: expected exactly one ID=5000 row, found {len(selected)}")
        template = selected.iloc[0]
        missing = sorted(set(columns).difference(template.index))
        if missing:
            raise ValueError(f"{process_id}: missing template columns {missing}")
        source_text = source.relative_to(PROJECT_ROOT).as_posix()
        template_cache[process_id] = (template, columns, source_text)

    # Process files contain a few case-only spelling differences such as
    # T_Cool1/T_COOL1.  Collapse them for portable wide CSV headers while the
    # long table retains the exact source variable spelling.
    source_condition_names = sorted(
        {name for _, columns, _ in template_cache.values() for name in columns},
        key=lambda value: (value.casefold(), value),
    )
    canonical_condition_name: dict[str, str] = {}
    for name in source_condition_names:
        canonical_condition_name.setdefault(name.casefold(), name)
    all_condition_names = sorted(canonical_condition_name.values(), key=str.casefold)

    for process_id in process_ids:
        template, columns, source_text = template_cache[process_id]
        baseline_runs = run_long.loc[run_long["process_id"] == process_id]
        baseline_condition_records.append(
            {
                "process_id": process_id,
                "baseline_definition": "unoptimized Process_Main ID=5000",
                "template_id": 5000,
                "template_source": source_text,
                "gnn_baseline_h2_mean_5seeds": baseline_runs["baseline_h2_production_kg_h"].mean(),
                "gnn_baseline_lcoh_mean_5seeds": baseline_runs["baseline_lcoh_usd_per_kg"].mean(),
                "gnn_baseline_specific_co2_mean_5seeds": baseline_runs[
                    "baseline_specific_co2_kg_per_kg_h2"
                ].mean(),
                **{
                    canonical_condition_name[name.casefold()]: as_float(template[name])
                    for name in columns
                },
            }
        )
        for name in columns:
            baseline_long_records.append(
                {
                    "process_id": process_id,
                    "template_id": 5000,
                    "template_source": source_text,
                    "variable_name": name,
                    "baseline_value": as_float(template[name]),
                    "unit_hint": unit_hint(name),
                }
            )

    queue_records: list[dict[str, Any]] = []
    condition_records: list[dict[str, Any]] = []
    condition_long_records: list[dict[str, Any]] = []
    candidate_result_records: list[dict[str, Any]] = []

    for row in candidates.sort_values(
        ["process_id", "process_rank_by_gnn_lcoh"],
        key=lambda series: series.map(process_number) if series.name == "process_id" else series,
    ).to_dict("records"):
        process_id = str(row["process_id"])
        run_id = str(row["run_id"])
        payload = payloads[run_id]
        process = process_payload(payload, process_id)
        baseline = process["baseline"]
        baseline_objectives = baseline["objectives"]
        baseline_indicators = baseline["indicators"]
        candidate = json.loads(str(row["candidate_json"]))
        template, columns, source_text = template_cache[process_id]
        unknown = sorted(set(candidate).difference(columns))
        if unknown:
            raise ValueError(f"{row['candidate_id']}: candidate variables not in safe columns: {unknown}")

        baseline_h2 = float(baseline_indicators["h2_production_kg_h"])
        baseline_lcoh = float(baseline_objectives["lcoh_usd_per_kg"])
        baseline_co2 = float(baseline_objectives["specific_co2_kg_per_kg_h2"])
        candidate_lcoh = float(row["gnn_lcoh_usd_per_kg"])
        lower_h2, upper_h2 = 0.9 * baseline_h2, 1.1 * baseline_h2
        prior = previous_by_process.get(process_id, {}) if row["aspen_validation_status"] == "completed" else {}
        prior_changes = as_float(prior.get("input_change_count")) if prior else None
        prior_h2 = as_float(prior.get("aspen_h2_production_kg_h")) if prior else None
        prior_lcoh = as_float(prior.get("aspen_lcoh_usd_per_kg")) if prior else None
        prior_co2 = as_float(prior.get("aspen_specific_co2_kg_per_kg_h2")) if prior else None

        base_record = {
            "candidate_id": row["candidate_id"],
            "process_id": process_id,
            "process_rank_by_gnn_lcoh": int(row["process_rank_by_gnn_lcoh"]),
            "seed": int(row["seed"]),
            "selected_rank_within_seed": int(row["selected_rank_within_seed"]),
            "source_run_id": run_id,
            "template_id": 5000,
            "template_source": source_text,
            "optimized_variable_count": len(candidate),
            "previous_validation_status": row["aspen_validation_status"],
            "gnn_baseline_h2_production_kg_h": baseline_h2,
            "gnn_baseline_lcoh_usd_per_kg": baseline_lcoh,
            "gnn_baseline_specific_co2_kg_per_kg_h2": baseline_co2,
            "gnn_candidate_h2_production_kg_h": float(row["gnn_h2_production_kg_h"]),
            "gnn_candidate_lcoh_usd_per_kg": candidate_lcoh,
            "gnn_candidate_specific_co2_kg_per_kg_h2": float(
                row["gnn_specific_co2_kg_per_kg_h2"]
            ),
            "gnn_lcoh_improvement_fraction": (baseline_lcoh - candidate_lcoh) / baseline_lcoh,
            "protocol_h2_lower_bound_kg_h": lower_h2,
            "protocol_h2_upper_bound_kg_h": upper_h2,
            "candidate_json": json.dumps(candidate, ensure_ascii=False, sort_keys=True),
        }
        queue_records.append(base_record)

        full_values = {name: as_float(candidate.get(name, template[name])) for name in columns}
        canonical_full_values = {
            canonical_condition_name[name.casefold()]: value
            for name, value in full_values.items()
        }
        condition_records.append(
            {
                **{key: base_record[key] for key in (
                    "candidate_id", "process_id", "process_rank_by_gnn_lcoh", "seed",
                    "selected_rank_within_seed", "source_run_id", "template_id",
                    "previous_validation_status",
                )},
                **{name: canonical_full_values.get(name) for name in all_condition_names},
            }
        )
        for name in columns:
            template_value = as_float(template[name])
            candidate_value = full_values[name]
            condition_long_records.append(
                {
                    "candidate_id": row["candidate_id"],
                    "process_id": process_id,
                    "process_rank_by_gnn_lcoh": int(row["process_rank_by_gnn_lcoh"]),
                    "seed": int(row["seed"]),
                    "variable_name": name,
                    "application_role": "optimized_by_ga" if name in candidate else "fixed_at_id5000",
                    "id5000_baseline_value": template_value,
                    "candidate_value": candidate_value,
                    "candidate_minus_baseline": (
                        candidate_value - template_value
                        if candidate_value is not None and template_value is not None else None
                    ),
                    "unit_hint": unit_hint(name),
                }
            )

        candidate_result_records.append(
            {
                **{key: base_record[key] for key in (
                    "candidate_id", "process_id", "process_rank_by_gnn_lcoh", "seed",
                    "selected_rank_within_seed", "source_run_id",
                    "gnn_baseline_h2_production_kg_h", "gnn_baseline_lcoh_usd_per_kg",
                    "gnn_baseline_specific_co2_kg_per_kg_h2",
                    "gnn_candidate_h2_production_kg_h", "gnn_candidate_lcoh_usd_per_kg",
                    "gnn_candidate_specific_co2_kg_per_kg_h2", "gnn_lcoh_improvement_fraction",
                    "protocol_h2_lower_bound_kg_h", "protocol_h2_upper_bound_kg_h",
                )},
                "aspen_run_status": "COMPLETED_REFERENCE" if prior else "PENDING",
                "aspen_case_file": "",
                "aspen_units_confirmed": "PENDING",
                "aspen_converged": "YES" if prior else "PENDING",
                "same_input_as_gnn_candidate": (
                    "YES" if prior and prior_changes == 0 else "NO" if prior else "PENDING"
                ),
                "input_change_count": prior_changes,
                "aspen_candidate_h2_production_kg_h": prior_h2,
                "aspen_candidate_lcoh_usd_per_kg": prior_lcoh,
                "aspen_candidate_specific_co2_kg_per_kg_h2": prior_co2,
                "aspen_baseline_h2_production_kg_h": None,
                "aspen_baseline_lcoh_usd_per_kg": None,
                "aspen_baseline_specific_co2_kg_per_kg_h2": None,
                "aspen_h2_protocol_window_pass": (
                    "YES" if prior_h2 is not None and lower_h2 <= prior_h2 <= upper_h2
                    else "NO" if prior_h2 is not None else "PENDING"
                ),
                "aspen_constraints_all_pass": "PENDING_RECHECK",
                "aspen_lcoh_improvement_fraction": None,
                "aspen_co2_change_fraction": None,
                "aspen_h2_change_vs_baseline_fraction": None,
                "gnn_vs_aspen_h2_relative_error_same_input_only": None,
                "gnn_vs_aspen_lcoh_relative_error_same_input_only": None,
                "gnn_vs_aspen_co2_relative_error_same_input_only": None,
                "final_candidate_decision": "PENDING_REVIEW",
                "aspen_notes": prior.get("validation_note", "") if prior else "",
            }
        )

    baseline_conditions = pd.DataFrame(baseline_condition_records)
    baseline_long = pd.DataFrame(baseline_long_records)
    queue = pd.DataFrame(queue_records)
    conditions = pd.DataFrame(condition_records)
    conditions_long = pd.DataFrame(condition_long_records)
    candidate_results = pd.DataFrame(candidate_result_records)

    baseline_results = pd.DataFrame(
        [
            {
                "process_id": process_id,
                "baseline_definition": "unoptimized Process_Main ID=5000",
                "template_id": 5000,
                "aspen_run_status": "PENDING",
                "aspen_case_file": "",
                "aspen_units_confirmed": "PENDING",
                "aspen_converged": "PENDING",
                "aspen_baseline_h2_production_kg_h": None,
                "aspen_baseline_lcoh_usd_per_kg": None,
                "aspen_baseline_specific_co2_kg_per_kg_h2": None,
                "aspen_constraints_all_pass": "PENDING",
                "aspen_notes": "",
            }
            for process_id in process_ids
        ]
    )

    criteria = pd.DataFrame(
        [
            ("Optimization baseline", "Same-process unoptimized Process_Main ID=5000 Aspen case"),
            ("Do not use", "Correction reference/Base row IDs such as 805 or 5937 are not improvement baselines"),
            ("LCOH improvement", "(Aspen baseline LCOH - Aspen candidate LCOH) / Aspen baseline LCOH"),
            ("CO2 change", "(Aspen candidate specific CO2 - Aspen baseline specific CO2) / Aspen baseline specific CO2"),
            ("H2 protocol window", "0.90 to 1.10 times the source-run GNN ID=5000 H2 value"),
            ("SC ratio", ">= 2.5"),
            ("Reformer outlet temperature", "<= 950 degC"),
            ("H2 purity", ">= 0.99 for P04/P06/P08/P09/P10 (P03 has no candidate)"),
            ("Minimum approach temperature", ">= 5 K"),
            ("Stream sanity", "nonnegative flow/fractions, fraction sum tolerance, max stream temperature <= 1300 degC"),
            ("Same-input accuracy", "Compare GNN and Aspen only when identical candidate inputs were used"),
            ("Pending label", "GNN-screened only; never report as Aspen-validated"),
        ],
        columns=["item", "requirement"],
    )
    return {
        "baseline_conditions": baseline_conditions,
        "baseline_long": baseline_long,
        "queue": queue,
        "candidate_conditions": conditions,
        "candidate_conditions_long": conditions_long,
        "baseline_results": baseline_results,
        "candidate_results": candidate_results,
        "previous_validation": previous,
        "criteria": criteria,
    }


def write_csvs(output_dir: Path, tables: dict[str, pd.DataFrame]) -> dict[str, Path]:
    names = {
        "baseline_conditions": "01_baseline_ID5000_conditions_wide.csv",
        "baseline_long": "02_baseline_ID5000_conditions_long.csv",
        "candidate_conditions": "03_candidate_conditions_wide.csv",
        "candidate_conditions_long": "04_candidate_conditions_long.csv",
        "baseline_results": "05_aspen_baseline_results_TO_FILL.csv",
        "candidate_results": "06_aspen_candidate_results_TO_FILL.csv",
        "queue": "07_gnn_screening_comparison.csv",
        "previous_validation": "08_previous_validation_reference.csv",
        "criteria": "09_validation_criteria.csv",
    }
    paths: dict[str, Path] = {}
    for key, filename in names.items():
        path = output_dir / filename
        tables[key].to_csv(path, index=False, encoding="utf-8-sig")
        paths[key] = path
    return paths


def add_excel_formulas(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    result_sheet = workbook["06_Aspen후보결과"]
    baseline_sheet = workbook["05_Aspen기준점결과"]
    headers = {cell.value: cell.column for cell in result_sheet[1]}
    baseline_headers = {cell.value: cell.column for cell in baseline_sheet[1]}

    process_col = get_column_letter(headers["process_id"])
    same_input_col = get_column_letter(headers["same_input_as_gnn_candidate"])
    gnn_h2_col = get_column_letter(headers["gnn_candidate_h2_production_kg_h"])
    gnn_lcoh_col = get_column_letter(headers["gnn_candidate_lcoh_usd_per_kg"])
    gnn_co2_col = get_column_letter(headers["gnn_candidate_specific_co2_kg_per_kg_h2"])
    aspen_h2_col = get_column_letter(headers["aspen_candidate_h2_production_kg_h"])
    aspen_lcoh_col = get_column_letter(headers["aspen_candidate_lcoh_usd_per_kg"])
    aspen_co2_col = get_column_letter(headers["aspen_candidate_specific_co2_kg_per_kg_h2"])

    baseline_lookup = {
        "aspen_baseline_h2_production_kg_h": "aspen_baseline_h2_production_kg_h",
        "aspen_baseline_lcoh_usd_per_kg": "aspen_baseline_lcoh_usd_per_kg",
        "aspen_baseline_specific_co2_kg_per_kg_h2": "aspen_baseline_specific_co2_kg_per_kg_h2",
    }
    for row_index in range(2, result_sheet.max_row + 1):
        for target_name, source_name in baseline_lookup.items():
            target_col = get_column_letter(headers[target_name])
            source_col = get_column_letter(baseline_headers[source_name])
            result_sheet[f"{target_col}{row_index}"] = (
                f'=IFERROR(INDEX(\'05_Aspen기준점결과\'!${source_col}:${source_col},'
                f'MATCH(${process_col}{row_index},\'05_Aspen기준점결과\'!$A:$A,0)),"")'
            )

        base_h2_col = get_column_letter(headers["aspen_baseline_h2_production_kg_h"])
        base_lcoh_col = get_column_letter(headers["aspen_baseline_lcoh_usd_per_kg"])
        base_co2_col = get_column_letter(headers["aspen_baseline_specific_co2_kg_per_kg_h2"])
        lower_h2_col = get_column_letter(headers["protocol_h2_lower_bound_kg_h"])
        upper_h2_col = get_column_letter(headers["protocol_h2_upper_bound_kg_h"])
        h2_pass_col = get_column_letter(headers["aspen_h2_protocol_window_pass"])
        improvement_col = get_column_letter(headers["aspen_lcoh_improvement_fraction"])
        co2_change_col = get_column_letter(headers["aspen_co2_change_fraction"])
        h2_change_col = get_column_letter(headers["aspen_h2_change_vs_baseline_fraction"])
        result_sheet[f"{h2_pass_col}{row_index}"] = (
            f'=IF({aspen_h2_col}{row_index}="","PENDING",'
            f'IF(AND({aspen_h2_col}{row_index}>={lower_h2_col}{row_index},'
            f'{aspen_h2_col}{row_index}<={upper_h2_col}{row_index}),"YES","NO"))'
        )
        result_sheet[f"{improvement_col}{row_index}"] = (
            f'=IFERROR(({base_lcoh_col}{row_index}-{aspen_lcoh_col}{row_index})/'
            f'{base_lcoh_col}{row_index},"")'
        )
        result_sheet[f"{co2_change_col}{row_index}"] = (
            f'=IFERROR(({aspen_co2_col}{row_index}-{base_co2_col}{row_index})/'
            f'{base_co2_col}{row_index},"")'
        )
        result_sheet[f"{h2_change_col}{row_index}"] = (
            f'=IFERROR(({aspen_h2_col}{row_index}-{base_h2_col}{row_index})/'
            f'{base_h2_col}{row_index},"")'
        )

        for target_name, gnn_col, aspen_col in (
            ("gnn_vs_aspen_h2_relative_error_same_input_only", gnn_h2_col, aspen_h2_col),
            ("gnn_vs_aspen_lcoh_relative_error_same_input_only", gnn_lcoh_col, aspen_lcoh_col),
            ("gnn_vs_aspen_co2_relative_error_same_input_only", gnn_co2_col, aspen_co2_col),
        ):
            target_col = get_column_letter(headers[target_name])
            result_sheet[f"{target_col}{row_index}"] = (
                f'=IF(${same_input_col}{row_index}<>"YES","",'
                f'IFERROR(ABS({gnn_col}{row_index}-{aspen_col}{row_index})/'
                f'MAX(ABS({aspen_col}{row_index}),1E-12),""))'
            )

    percent_headers = {
        "gnn_lcoh_improvement_fraction",
        "aspen_lcoh_improvement_fraction",
        "aspen_co2_change_fraction",
        "aspen_h2_change_vs_baseline_fraction",
        "gnn_vs_aspen_h2_relative_error_same_input_only",
        "gnn_vs_aspen_lcoh_relative_error_same_input_only",
        "gnn_vs_aspen_co2_relative_error_same_input_only",
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        local_headers = {cell.value: cell.column for cell in sheet[1]}
        for header in percent_headers.intersection(local_headers):
            for cell in sheet.iter_cols(
                min_col=local_headers[header], max_col=local_headers[header], min_row=2
            ):
                for item in cell:
                    item.number_format = "0.00%"
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            sample = list(column_cells[:200])
            width = max((len(str(cell.value)) for cell in sample if cell.value is not None), default=8)
            sheet.column_dimensions[letter].width = min(max(width + 2, 10), 42)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(workbook_path)


def write_workbook(output_dir: Path, tables: dict[str, pd.DataFrame]) -> Path:
    workbook_path = output_dir / "Aspen_GA_Validation_Handoff_135.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        tables["criteria"].to_excel(writer, sheet_name="00_필독", index=False)
        tables["baseline_conditions"].to_excel(writer, sheet_name="01_ID5000기준입력", index=False)
        tables["candidate_conditions"].to_excel(writer, sheet_name="02_후보전체입력", index=False)
        tables["candidate_conditions_long"].to_excel(writer, sheet_name="03_후보입력long", index=False)
        tables["queue"].to_excel(writer, sheet_name="04_GNN선별결과", index=False)
        tables["baseline_results"].to_excel(writer, sheet_name="05_Aspen기준점결과", index=False)
        tables["candidate_results"].to_excel(writer, sheet_name="06_Aspen후보결과", index=False)
        tables["previous_validation"].to_excel(writer, sheet_name="07_기존9개참고", index=False)
    add_excel_formulas(workbook_path)
    return workbook_path


def readme_text(candidate_count: int, completed_count: int) -> str:
    pending_count = candidate_count - completed_count
    return f"""# Aspen GA 후보 확대 검증 전달 문서

## 1. 목적

이 패키지는 `LCOH_capacity / GA`에서 선별한 **{candidate_count}개 후보**를 Aspen으로
독립 검증하고, 각 공정의 기존 운전점 대비 LCOH 개선률을 계산하기 위한 자료입니다.

- 대상 공정: P01, P02, P04–P10 (9개 공정)
- 공정별 후보: 15개
- 선정 방식: 5개 seed 각각에서 feasible·중복 제거 후 GNN LCOH 상위 3개
- 기존 Aspen 참고 결과: {completed_count}개
- 추가 검증 대상: {pending_count}개
- P03: GA feasible 후보가 없어 제외

## 2. 가장 중요한 기준점

개선률 기준점은 각 공정 `Process_Main`의 **ID=5000 원 운전점**입니다.

`Aspen_GNN_Optimization_Validation_20260915_revised.xlsx`에 기록된 `Base 행 ID`
(예: 805, 5937)는 수렴 보정용 최근접 데이터 행입니다. 개선률 기준점이 아니므로
절대 사용하지 마십시오.

Aspen 검증 개선률은 반드시 Aspen끼리 계산합니다.

```text
Aspen LCOH 개선률 =
    (ID=5000 Aspen LCOH - 후보 Aspen LCOH) / ID=5000 Aspen LCOH
```

GNN 기준점과 Aspen 후보를 섞어서 개선률을 계산하지 마십시오.

## 3. 권장 실행 순서

1. `01_baseline_ID5000_conditions_wide.csv`를 사용해 9개 공정의 ID=5000 기준점을
   Aspen에서 먼저 실행합니다.
2. 결과를 `05_aspen_baseline_results_TO_FILL.csv` 또는 Excel의
   `05_Aspen기준점결과` 시트에 기록합니다.
3. `03_candidate_conditions_wide.csv`의 후보를 Aspen에 적용합니다. 후보 JSON에
   없는 안전 입력은 모두 ID=5000 값으로 고정되어 있습니다.
4. Reinit 후 Run2를 실행하고 UOSSTAT, block error, 단위, 수렴 여부를 기록합니다.
5. 후보 결과를 `06_aspen_candidate_results_TO_FILL.csv` 또는 Excel의
   `06_Aspen후보결과` 시트에 기록합니다.
6. Excel을 열면 ID=5000 Aspen 결과를 공정별로 연결해 개선률이 자동 계산됩니다.

## 4. 원 입력과 보정 입력

- 먼저 제공된 후보 입력을 수정하지 않고 실행하고 그 결과를 보존하십시오.
- 미수렴하여 물리 연동 제약을 보정해야 하면 원 입력 결과와 보정 결과를 분리합니다.
- 변경 변수, 변경 전후 값, 변경 이유를 반드시 기록합니다.
- 보정 입력에서 GNN–Aspen 오차를 보고하려면 보정된 동일 입력으로 GNN을 다시
  추론해야 합니다. 원 후보 GNN 값과 보정 Aspen 값을 동일 입력 정확도처럼 비교하면 안 됩니다.
- `same_input_as_gnn_candidate=YES`인 경우에만 Excel의 GNN–Aspen 오차식을 사용합니다.

## 5. 제약 판정

후보가 수렴했다는 사실만으로 최적해가 되는 것은 아닙니다. 다음을 모두 확인하십시오.

- H2 생산량: source-run GNN ID=5000 생산량의 90–110% 숫자 범위
- S/C ratio >= 2.5
- reformer outlet temperature <= 950 degC
- P04/P06/P08/P09/P10 제품 H2 purity >= 0.99
- minimum approach temperature >= 5 K
- 음의 유량/몰분율, 몰분율 합, 최고 스트림 온도 등 stream sanity

`aspen_constraints_all_pass=YES`이고 정상수렴한 후보만 최종 Aspen-validated
최적 후보로 사용할 수 있습니다.

## 6. 파일 안내

- `Aspen_GA_Validation_Handoff_135.xlsx`: 통합 작업 파일 및 자동 개선률 계산
- `01_baseline_ID5000_conditions_wide.csv`: 공정별 원 운전점 전체 입력
- `02_baseline_ID5000_conditions_long.csv`: 기준점 입력 long 형식
- `03_candidate_conditions_wide.csv`: 135개 후보 전체 적용 입력
- `04_candidate_conditions_long.csv`: 후보별 변수·ID=5000 대비 변경값
- `05_aspen_baseline_results_TO_FILL.csv`: 기준점 Aspen 결과 입력용
- `06_aspen_candidate_results_TO_FILL.csv`: 후보 Aspen 결과 입력용
- `07_gnn_screening_comparison.csv`: GNN 기준점·후보·예상 개선률
- `08_previous_validation_reference.csv`: 기존 9개 Aspen 결과(참고용)
- `09_validation_criteria.csv`: 판정 기준
- `provenance/`: 원 후보와 기존 검증 파일 및 출처

## 7. 결과 회신 시 필요한 파일

최소한 다음 두 시트 또는 CSV를 채워 회신해 주십시오.

1. `05_Aspen기준점결과`
2. `06_Aspen후보결과`

Aspen case 파일명, 수렴 여부, 적용 입력 변경 수, H2 생산량, LCOH, specific CO2,
전체 제약 통과 여부와 메모를 함께 남겨야 합니다.

## 8. 결과 표현

- GNN 결과: `GNN-screened improvement`
- Aspen 결과: ID=5000 기준점과 후보가 모두 Aspen 계산되고 제약을 통과한 경우에만
  `Aspen-validated improvement`
- GA 알고리즘 성능 비교는 별도로 동일 예산 random search와 비교합니다.
"""


def main() -> None:
    args = parse_args()
    study_root = resolve_path(args.study_root)
    candidate_dir = resolve_path(args.candidate_dir)
    validation_workbook = resolve_path(args.validation_workbook)
    output_dir = resolve_path(args.output_dir)
    if output_dir.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output_dir}")

    candidates_path = candidate_dir / "expanded_aspen_candidates.csv"
    candidates = pd.read_csv(candidates_path)
    if len(candidates) < 100:
        raise ValueError(f"Expected at least 100 candidates, found {len(candidates)}")
    if candidates.duplicated("candidate_id").any():
        raise ValueError("Duplicate candidate_id values")
    payloads = load_payloads(study_root, candidates)
    previous = load_previous_validation(validation_workbook)
    tables = build_tables(study_root, candidates, payloads, previous)

    output_dir.mkdir(parents=True, exist_ok=False)
    provenance_dir = output_dir / "provenance"
    provenance_dir.mkdir()
    csv_paths = write_csvs(output_dir, tables)
    workbook_path = write_workbook(output_dir, tables)
    completed_count = int((candidates["aspen_validation_status"] == "completed").sum())
    (output_dir / "README.md").write_text(
        readme_text(len(candidates), completed_count), encoding="utf-8"
    )
    (output_dir / "전달메시지.txt").write_text(
        "안녕하세요. GA로 선별한 LCOH 최적화 후보의 Aspen 검증을 부탁드립니다.\n\n"
        "후보는 P03을 제외한 9개 공정, 총 135개입니다. 개선률 기준점은 각 공정의 "
        "Process_Main ID=5000 원 운전점입니다. 먼저 기준점 9개를 Aspen으로 계산한 뒤 "
        "후보를 검증해 주세요. Excel의 05_Aspen기준점결과와 06_Aspen후보결과 시트에 "
        "수렴 여부, H2, LCOH, specific CO2, 입력 보정 여부와 제약 통과 여부를 기록해 "
        "주시면 개선률이 자동 계산됩니다. 자세한 절차와 주의사항은 README.md를 먼저 "
        "확인해 주세요. 특히 기존 파일의 Base 행 ID는 개선률 기준점이 아닙니다.\n",
        encoding="utf-8",
    )

    source_files = [
        candidates_path,
        candidate_dir / "selection_summary.csv",
        candidate_dir / "manifest.json",
        validation_workbook,
        study_root / "paper_summary" / "paper_report.json",
        study_root / "paper_summary" / "paper_runs_long.csv",
    ]
    for source in source_files:
        if source.is_file():
            shutil.copy2(source, provenance_dir / source.name)

    manifest = {
        "purpose": "Expanded GA candidate Aspen validation and ID=5000 improvement calculation",
        "candidate_count": int(len(candidates)),
        "process_count": int(candidates["process_id"].nunique()),
        "baseline_case_count": int(len(tables["baseline_results"])),
        "previously_completed_reference_count": completed_count,
        "pending_candidate_count": int(len(candidates) - completed_count),
        "baseline_definition": "same-process Process_Main ID=5000",
        "workbook": workbook_path.name,
        "csv_files": {key: path.name for key, path in csv_paths.items()},
        "files": [],
    }
    for path in sorted(output_dir.rglob("*")):
        if path.is_file() and path.name != "manifest.json":
            manifest["files"].append(
                {
                    "path": path.relative_to(output_dir).as_posix(),
                    "sha256": sha256(path),
                    "bytes": path.stat().st_size,
                }
            )
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(
        f"[DELIVERY] output={output_dir}; candidates={len(candidates)}; "
        f"baselines={len(tables['baseline_results'])}; completed_reference={completed_count}"
    )


if __name__ == "__main__":
    main()
