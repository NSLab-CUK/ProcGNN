"""Select a balanced, expanded Aspen-validation queue from GA histories.

The paper summarizer retains only one best feasible point per seed.  This
utility reads candidate-level ``ga_history.csv`` files so that additional
high-quality points can be selected without rerunning the optimisation.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--study-root", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--scenario", default="LCOH_capacity")
    parser.add_argument("--method", default="ga")
    parser.add_argument("--per-seed", type=int, default=3)
    parser.add_argument("--expected-minimum", type=int, default=100)
    parser.add_argument("--validation-workbook", default=None)
    return parser.parse_args()


def resolve_path(value: str) -> Path:
    path = Path(value)
    return path.resolve() if path.is_absolute() else (PROJECT_ROOT / path).resolve()


def canonical_candidate(value: object) -> tuple[str, dict[str, float]]:
    payload = json.loads(str(value))
    if not isinstance(payload, dict):
        raise TypeError("candidate_json must decode to an object")
    numeric = {str(key): float(item) for key, item in payload.items()}
    canonical = json.dumps(numeric, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return canonical, numeric


def relative_error(prediction: object, reference: object) -> float | None:
    if prediction is None or reference is None:
        return None
    prediction_f = float(prediction)
    reference_f = float(reference)
    if abs(prediction_f) < 1e-8 and abs(reference_f) < 1e-8:
        return 0.0
    return abs(prediction_f - reference_f) / max(abs(reference_f), 1e-12)


def load_previous_queue(study_root: Path, scenario: str) -> dict[tuple[str, str], int]:
    queue_path = study_root / "paper_summary" / "aspen_recheck_candidates.csv"
    if not queue_path.is_file():
        return {}
    queue = pd.read_csv(queue_path)
    queue = queue.loc[queue["scenario"].astype(str) == scenario]
    result: dict[tuple[str, str], int] = {}
    for row in queue.to_dict("records"):
        canonical, _ = canonical_candidate(row["candidate_json"])
        result[(str(row["process_id"]), canonical)] = int(row["rank_by_gnn_lcoh"])
    return result


def load_previous_validation(workbook_path: Path | None) -> dict[tuple[str, str], dict[str, Any]]:
    if workbook_path is None:
        return {}
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    candidates = workbook["전체후보목록"]
    validated_candidates: dict[tuple[str, str], bool] = {}
    for row_index in range(5, candidates.max_row + 1):
        if str(candidates.cell(row_index, 1).value or "").strip() != "예":
            continue
        process_id = str(candidates.cell(row_index, 3).value)
        canonical, _ = canonical_candidate(candidates.cell(row_index, 11).value)
        validated_candidates[(process_id, canonical)] = True

    summary = workbook["결과요약"]
    by_process: dict[str, dict[str, Any]] = {}
    for row_index in range(9, 19):
        process_id = summary.cell(row_index, 1).value
        if not process_id:
            continue
        by_process[str(process_id)] = {
            "previous_aspen_status": summary.cell(row_index, 3).value,
            "previous_input_change_count": summary.cell(row_index, 6).value,
            "previous_h2_relative_error": relative_error(
                summary.cell(row_index, 8).value, summary.cell(row_index, 9).value
            ),
            "previous_lcoh_relative_error": relative_error(
                summary.cell(row_index, 12).value, summary.cell(row_index, 13).value
            ),
            "previous_co2_relative_error": relative_error(
                summary.cell(row_index, 16).value, summary.cell(row_index, 17).value
            ),
            "previous_validation_note": summary.cell(row_index, 7).value,
        }

    result: dict[tuple[str, str], dict[str, Any]] = {}
    for key in validated_candidates:
        result[key] = by_process.get(key[0], {})
    return result


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_histories(
    study_root: Path, scenario: str, method: str
) -> tuple[pd.DataFrame, list[Path]]:
    history_paths = sorted((study_root / scenario / method).glob("seed_*/ga_history.csv"))
    if not history_paths:
        raise FileNotFoundError(f"No histories found under {study_root / scenario / method}")

    frames: list[pd.DataFrame] = []
    required = {
        "candidate_json",
        "feasible",
        "generation",
        "member",
        "process_id",
        "h2_production_kg_h",
        "lcoh_usd_per_kg",
        "specific_co2_kg_per_kg_h2",
        "violation_sum",
    }
    for path in history_paths:
        frame = pd.read_csv(path)
        missing = required.difference(frame.columns)
        if missing:
            raise ValueError(f"{path} is missing columns: {sorted(missing)}")
        seed_text = path.parent.name.removeprefix("seed_")
        frame["seed"] = int(seed_text)
        frame["run_id"] = f"{scenario}/{method}/seed_{seed_text}"
        frames.append(frame)

    history = pd.concat(frames, ignore_index=True)
    history["feasible"] = history["feasible"].astype(str).str.lower().eq("true")
    for column in (
        "generation",
        "member",
        "h2_production_kg_h",
        "lcoh_usd_per_kg",
        "specific_co2_kg_per_kg_h2",
        "violation_sum",
    ):
        history[column] = pd.to_numeric(history[column], errors="raise")

    canonical_values = history["candidate_json"].map(canonical_candidate)
    history["candidate_key"] = canonical_values.map(lambda item: item[0])
    history["candidate_values"] = canonical_values.map(lambda item: item[1])
    return history, history_paths


def select_candidates(
    history: pd.DataFrame,
    scenario: str,
    method: str,
    per_seed: int,
    previous_queue: dict[tuple[str, str], int],
    previous_validation: dict[tuple[str, str], dict[str, Any]],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    feasible = history.loc[history["feasible"]].copy()
    process_ids = sorted(history["process_id"].astype(str).unique())
    seeds = sorted(int(value) for value in history["seed"].unique())
    selected_records: list[dict[str, Any]] = []
    summary_records: list[dict[str, Any]] = []

    for process_id in process_ids:
        process_feasible = feasible.loc[feasible["process_id"].astype(str) == process_id].copy()
        process_feasible = process_feasible.sort_values(
            ["lcoh_usd_per_kg", "seed", "generation", "member"], kind="stable"
        )
        process_unique = process_feasible.drop_duplicates("candidate_key", keep="first")
        seen: set[str] = set()
        process_selected: list[dict[str, Any]] = []

        for seed in seeds:
            seed_rows = process_feasible.loc[process_feasible["seed"] == seed]
            seed_rows = seed_rows.drop_duplicates("candidate_key", keep="first").reset_index(drop=True)
            chosen_for_seed = 0
            for source_index, row in seed_rows.iterrows():
                candidate_key = str(row["candidate_key"])
                if candidate_key in seen:
                    continue
                chosen_for_seed += 1
                seen.add(candidate_key)
                record = row.to_dict()
                record["selected_rank_within_seed"] = chosen_for_seed
                record["source_feasible_rank_within_seed"] = source_index + 1
                process_selected.append(record)
                if chosen_for_seed == per_seed:
                    break
            if len(seed_rows) and chosen_for_seed < per_seed:
                raise RuntimeError(
                    f"{process_id}/seed_{seed}: selected only {chosen_for_seed} of {per_seed}"
                )

        process_selected.sort(
            key=lambda row: (
                float(row["lcoh_usd_per_kg"]),
                int(row["seed"]),
                int(row["generation"]),
                int(row["member"]),
            )
        )
        for process_rank, row in enumerate(process_selected, start=1):
            candidate_key = str(row["candidate_key"])
            prior = previous_validation.get((process_id, candidate_key), {})
            seed = int(row["seed"])
            seed_rank = int(row["selected_rank_within_seed"])
            selected_records.append(
                {
                    "candidate_id": f"{scenario}_{process_id}_S{seed}_SR{seed_rank:02d}",
                    "scenario": scenario,
                    "search_method": method,
                    "process_id": process_id,
                    "process_rank_by_gnn_lcoh": process_rank,
                    "seed": seed,
                    "selected_rank_within_seed": seed_rank,
                    "source_feasible_rank_within_seed": int(
                        row["source_feasible_rank_within_seed"]
                    ),
                    "run_id": row["run_id"],
                    "generation": int(row["generation"]),
                    "member": int(row["member"]),
                    "gnn_h2_production_kg_h": float(row["h2_production_kg_h"]),
                    "gnn_lcoh_usd_per_kg": float(row["lcoh_usd_per_kg"]),
                    "gnn_specific_co2_kg_per_kg_h2": float(
                        row["specific_co2_kg_per_kg_h2"]
                    ),
                    "violation_sum": float(row["violation_sum"]),
                    "existing_queue_rank": previous_queue.get((process_id, candidate_key)),
                    "aspen_validation_status": "completed" if prior else "pending",
                    "previous_aspen_status": prior.get("previous_aspen_status"),
                    "previous_input_change_count": prior.get("previous_input_change_count"),
                    "previous_h2_relative_error": prior.get("previous_h2_relative_error"),
                    "previous_lcoh_relative_error": prior.get("previous_lcoh_relative_error"),
                    "previous_co2_relative_error": prior.get("previous_co2_relative_error"),
                    "previous_validation_note": prior.get("previous_validation_note"),
                    "candidate_json": json.dumps(
                        row["candidate_values"], ensure_ascii=False, sort_keys=True
                    ),
                    "candidate_key": candidate_key,
                    "candidate_values": row["candidate_values"],
                }
            )

        summary_records.append(
            {
                "process_id": process_id,
                "source_feasible_rows": int(len(process_feasible)),
                "source_distinct_feasible_candidates": int(len(process_unique)),
                "selected_candidates": int(len(process_selected)),
                "seed_count": int(len({int(row["seed"]) for row in process_selected})),
                "selection_status": "selected" if process_selected else "no_feasible_candidate",
            }
        )

    selected = pd.DataFrame(selected_records)
    summary = pd.DataFrame(summary_records)
    return selected, summary


def candidate_tables(selected: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    variable_names = sorted(
        {name for payload in selected["candidate_values"] for name in payload}
    )
    wide_rows: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    for row in selected.to_dict("records"):
        base = {
            "candidate_id": row["candidate_id"],
            "process_id": row["process_id"],
            "process_rank_by_gnn_lcoh": row["process_rank_by_gnn_lcoh"],
            "seed": row["seed"],
            "selected_rank_within_seed": row["selected_rank_within_seed"],
            "aspen_validation_status": row["aspen_validation_status"],
        }
        payload = row["candidate_values"]
        wide_rows.append({**base, **{name: payload.get(name) for name in variable_names}})
        for name, value in sorted(payload.items()):
            long_rows.append({**base, "variable_name": name, "candidate_value": value})
    return pd.DataFrame(wide_rows), pd.DataFrame(long_rows)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            sample = list(column_cells[:200])
            max_length = max((len(str(cell.value)) for cell in sample if cell.value is not None), default=8)
            sheet.column_dimensions[letter].width = min(max(max_length + 2, 10), 45)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    if args.per_seed < 1:
        raise ValueError("--per-seed must be >= 1")
    if args.expected_minimum < 1:
        raise ValueError("--expected-minimum must be >= 1")

    study_root = resolve_path(args.study_root)
    output_dir = resolve_path(args.output_dir)
    if output_dir.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output_dir}")
    validation_workbook = (
        resolve_path(args.validation_workbook) if args.validation_workbook else None
    )

    history, history_paths = read_histories(study_root, args.scenario, args.method)
    previous_queue = load_previous_queue(study_root, args.scenario)
    previous_validation = load_previous_validation(validation_workbook)
    selected, summary = select_candidates(
        history,
        scenario=args.scenario,
        method=args.method,
        per_seed=args.per_seed,
        previous_queue=previous_queue,
        previous_validation=previous_validation,
    )
    if len(selected) < args.expected_minimum:
        raise RuntimeError(
            f"Selected {len(selected)} candidates, below required minimum {args.expected_minimum}"
        )
    if selected.duplicated(["process_id", "candidate_key"]).any():
        raise RuntimeError("Duplicate candidates remain within a process")

    wide, long = candidate_tables(selected)
    export_columns = [column for column in selected.columns if column not in {"candidate_key", "candidate_values"}]
    exported = selected[export_columns].copy()
    output_dir.mkdir(parents=True, exist_ok=False)
    exported.to_csv(output_dir / "expanded_aspen_candidates.csv", index=False, encoding="utf-8-sig")
    summary.to_csv(output_dir / "selection_summary.csv", index=False, encoding="utf-8-sig")
    wide.to_csv(output_dir / "candidate_operating_conditions_wide.csv", index=False, encoding="utf-8-sig")
    long.to_csv(output_dir / "candidate_variables_long.csv", index=False, encoding="utf-8-sig")

    criteria = pd.DataFrame(
        [
            ("study_root", str(study_root)),
            ("scenario", args.scenario),
            ("search_method", args.method),
            ("selection", "feasible and exact-deduplicated, lowest GNN LCOH within each process/seed"),
            ("per_seed", args.per_seed),
            ("selected_count", len(selected)),
            ("expected_minimum", args.expected_minimum),
            ("important", "pending rows are GNN-screened candidates and are not Aspen-validated"),
        ],
        columns=["item", "value"],
    )
    workbook_path = output_dir / "Expanded_Aspen_Validation_Candidates.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="선정요약", index=False)
        exported.to_excel(writer, sheet_name="후보목록", index=False)
        wide.to_excel(writer, sheet_name="운전변수_wide", index=False)
        long.to_excel(writer, sheet_name="운전변수_long", index=False)
        criteria.to_excel(writer, sheet_name="선정기준", index=False)
    format_workbook(workbook_path)

    manifest = {
        "study_root": str(study_root),
        "scenario": args.scenario,
        "search_method": args.method,
        "seeds": sorted(int(value) for value in selected["seed"].unique()),
        "per_seed": args.per_seed,
        "selected_candidates": int(len(selected)),
        "processes_with_candidates": sorted(selected["process_id"].astype(str).unique()),
        "completed_previous_aspen_validation": int(
            (selected["aspen_validation_status"] == "completed").sum()
        ),
        "pending_aspen_validation": int(
            (selected["aspen_validation_status"] == "pending").sum()
        ),
        "source_histories": [
            {"path": str(path), "sha256": sha256(path)} for path in history_paths
        ],
        "validation_workbook": str(validation_workbook) if validation_workbook else None,
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    (output_dir / "README.md").write_text(
        f"""# Expanded Aspen validation candidates

- Scope: `{args.scenario}` / `{args.method}`
- Rule: feasible candidates only; exact candidate JSON deduplicated within each process;
  select the lowest-LCOH `{args.per_seed}` candidates from each seed.
- Selected: **{len(selected)}** candidates.
- Previously Aspen-validated: **{manifest['completed_previous_aspen_validation']}**.
- Pending Aspen validation: **{manifest['pending_aspen_validation']}**.
- P03 remains absent because the source GA histories contain no feasible candidate.

`pending` means GNN-screened only.  Do not report a pending row as Aspen-validated.
The workbook records the original candidate operating variables; physical coupling
constraints and same-input GNN re-inference must be handled during Aspen validation.
""",
        encoding="utf-8",
    )
    print(
        f"[EXPANDED ASPEN QUEUE] output={output_dir}; selected={len(selected)}; "
        f"completed={manifest['completed_previous_aspen_validation']}; "
        f"pending={manifest['pending_aspen_validation']}"
    )


if __name__ == "__main__":
    main()
