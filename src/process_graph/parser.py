from __future__ import annotations

import re
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .constants import (
    DECODER_CATEGORY_MAP,
    OPER_FEATURE_SLOTS,
    UNIT_TYPE_TO_CANONICAL,
)
from .schema import CellSpec, NodeSpec, ProcessSpec

# User-approved fixed mapping by process id:
# (target H2 node, tailgas CO2 node).
PROCESS_FIXED_GAS_READOUT: Dict[str, Tuple[str, str]] = {
    "Process1": ("OUT_prod", "OUT_exhaust"),
    "Process2": ("OUT_PROD", "OUT_EXHAUST"),
    "Process3": ("OUT_PROD", "OUT_CO2"),
    "Process4": ("OUT_PROD", "OUT_EXHAUST"),
    "Process5": ("OUT_PROD", "OUT_CO2"),
    "Process6": ("OUT_PROD", "OUT_CO2"),
    "Process7": ("OUT_PROD", "OUT_PROD"),
    "Process8": ("OUT_PROD", "OUT_CO2"),
    "Process9": ("OUT_PROD", "OUT_CO2"),
    "Process10": ("OUT_PROD", "OUT_CO2"),
}


def _clean_str(value) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def infer_node_role(node_name: str) -> str:
    name = node_name.upper()
    if name.startswith("IN_") or name.startswith("INPUT_"):
        return "source"
    if name.startswith("OUT_"):
        return "sink"
    return "unit"


def canonicalize_type(type_raw: str) -> str:
    if type_raw not in UNIT_TYPE_TO_CANONICAL:
        raise ValueError(f"Unknown unit type: {type_raw}")
    return UNIT_TYPE_TO_CANONICAL[type_raw]


def infer_hx_role(type_raw: str):
    if type_raw == "HX_hot":
        return "hot_outlet"
    if type_raw == "HX_cold":
        return "cold_outlet"
    if type_raw == "HX_dt":
        return "hot_inlet_cold_outlet_diff"
    return None


def parse_cell_spec(value) -> CellSpec:
    if value is None:
        return CellSpec(kind="missing", value=None)

    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return CellSpec(kind="missing", value=None)
        try:
            num = float(s)
            return CellSpec(kind="constant", value=num)
        except ValueError:
            return CellSpec(kind="reference", value=s)

    if isinstance(value, (int, float)):
        return CellSpec(kind="constant", value=float(value))

    raise TypeError(f"Unsupported cell value: {value}")


def parse_burner_pressure_cell(value) -> CellSpec:
    if value is None:
        return CellSpec(kind="missing", value=None)

    if isinstance(value, (int, float)):
        num = float(value)
        if num == 0.0:
            return CellSpec(kind="passthrough", value=None)
        return CellSpec(kind="constant", value=num)

    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return CellSpec(kind="missing", value=None)

        if s == "0":
            return CellSpec(kind="passthrough", value=None)

        m = re.match(r"^0\s*\((.+)\)$", s)
        if m:
            hint_ref = m.group(1).strip()
            return CellSpec(kind="passthrough", value=None, meta={"hint_ref": hint_ref})

        try:
            num = float(s)
            if num == 0.0:
                return CellSpec(kind="passthrough", value=None)
            return CellSpec(kind="constant", value=num)
        except ValueError:
            return CellSpec(kind="reference", value=s)

    raise TypeError(f"Unsupported burner pressure cell: {value}")


def is_adjacency_sheet(ws: Worksheet) -> bool:
    if ws.max_row < 3 or ws.max_column < 3:
        return False

    top = []
    for c in range(2, ws.max_column + 1):
        val = _clean_str(ws.cell(1, c).value)
        if val is not None:
            top.append(val)

    left = []
    for r in range(2, ws.max_row + 1):
        val = _clean_str(ws.cell(r, 1).value)
        if val is not None:
            left.append(val)

    if len(top) < 2 or len(left) < 2:
        return False
    if top != left:
        return False

    observed = []
    check_rows = min(len(left), 10)
    check_cols = min(len(top), 10)
    for r in range(2, 2 + check_rows):
        for c in range(2, 2 + check_cols):
            observed.append(ws.cell(r, c).value)

    valid_count = 0
    for v in observed:
        if v in (0, 1, 0.0, 1.0, None, ""):
            valid_count += 1

    return valid_count / max(len(observed), 1) >= 0.8


def is_node_spec_sheet(ws: Worksheet) -> bool:
    found_type = False
    for r in range(1, min(ws.max_row, 4) + 1):
        for c in range(1, min(ws.max_column, 5) + 1):
            if _clean_str(ws.cell(r, c).value) == "Type":
                found_type = True
                break
        if found_type:
            break
    if not found_type:
        return False

    slot_hits = 0
    for c in range(1, ws.max_column + 1):
        val = _clean_str(ws.cell(2, c).value)
        if val in OPER_FEATURE_SLOTS:
            slot_hits += 1

    return slot_hits >= 4


def is_decoder_sheet(ws: Worksheet) -> bool:
    targets = set(DECODER_CATEGORY_MAP.keys())
    hits = 0
    for r in range(1, min(ws.max_row, 5) + 1):
        for c in range(1, ws.max_column + 1):
            val = _clean_str(ws.cell(r, c).value)
            if val in targets:
                hits += 1
    return hits >= 2


def detect_sheet_roles(workbook) -> Dict[str, Worksheet]:
    adjacency_ws = None
    node_ws = None
    decoder_ws = None

    for ws in workbook.worksheets:
        if adjacency_ws is None and is_adjacency_sheet(ws):
            adjacency_ws = ws
            continue
        if node_ws is None and is_node_spec_sheet(ws):
            node_ws = ws
            continue
        if decoder_ws is None and is_decoder_sheet(ws):
            decoder_ws = ws
            continue

    if adjacency_ws is None or node_ws is None or decoder_ws is None:
        found = {
            "adjacency": adjacency_ws.title if adjacency_ws else None,
            "node_spec": node_ws.title if node_ws else None,
            "decoder": decoder_ws.title if decoder_ws else None,
        }
        raise ValueError(f"Failed to detect required sheets correctly: {found}")

    return {
        "adjacency": adjacency_ws,
        "node_spec": node_ws,
        "decoder": decoder_ws,
    }


def _cell_fill_rgb_hex(cell) -> Optional[str]:
    """Return ARGB hex for solid fill foreground, or None."""
    fill = cell.fill
    if fill is None or fill.patternType != "solid" or fill.fgColor is None:
        return None
    color = fill.fgColor
    if color.type == "rgb" and color.rgb:
        return str(color.rgb).upper()
    return None


def _is_yellow_highlight_cell(cell) -> bool:
    """True for standard yellow highlight (Excel ARGB FFFFFF00)."""
    rgb = _cell_fill_rgb_hex(cell)
    if rgb is None:
        return False
    return rgb in {"FFFFFF00", "FFFF00", "00FFFF00"}


def collect_yellow_sink_names_second_sheet(workbook) -> List[str]:
    """
    Read column A on the workbook's **second worksheet** (index 1).

    Convention (per project workbooks): output sinks are marked with yellow fill
    on the node name cell in column A. Order is top-to-bottom as in the sheet.
    """
    if len(workbook.worksheets) < 2:
        return []
    ws = workbook.worksheets[1]
    names: List[str] = []
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, 1)
        if not _is_yellow_highlight_cell(cell):
            continue
        name = _clean_str(cell.value)
        if name is None:
            continue
        names.append(name)
    return names


def _match_graph_node_name(raw: str, node_names: List[str]) -> Optional[str]:
    """Map a sheet string to the exact name used in ``node_names``."""
    if raw in node_names:
        return raw
    raw_fold = raw.casefold()
    for nm in node_names:
        if nm.casefold() == raw_fold:
            return nm
    return None


def resolve_gas_readout_nodes_from_highlight_list(
    highlighted_sink_names: List[str],
    node_names: List[str],
) -> Dict[str, str]:
    """
    From yellow-marked output sinks, pick readout nodes for fixed H2 / CO2 tasks.

    - ``target`` (PROD_H2_MoleFlow, etc.): primary product / hydrogen outlet.
    - ``tailgas`` (PROD_CO2_MoleFlow, etc.): dedicated CO2 outlet, else exhaust / tail stream.

    Uses sheet order as tie-breaker among equally-scored options.
    """
    ordered: List[str] = []
    for raw in highlighted_sink_names:
        m = _match_graph_node_name(raw, node_names)
        if m is not None and m not in ordered:
            ordered.append(m)
    if not ordered:
        return {}

    def pick_h2() -> Optional[str]:
        for m in ordered:
            u = m.upper()
            if u in ("OUT_PROD", "OUT_PRODUCT") or (u.endswith("_PROD") and "CO2" not in u):
                return m
        for m in ordered:
            u = m.upper()
            if "PROD" in u and "CO2" not in u and "EXHAUST" not in u and "STREAM" not in u:
                return m
        for m in ordered:
            if "PROD" in m.upper():
                return m
        for m in ordered:
            u = m.upper()
            if u.startswith("OUT_") and "CO2" not in u:
                return m
        return ordered[0]

    def pick_co2() -> Optional[str]:
        for m in ordered:
            if "CO2" in m.upper():
                return m
        for m in ordered:
            if "EXHAUST" in m.upper():
                return m
        for m in ordered:
            if "STREAM" in m.upper():
                return m
        for m in ordered:
            if "H2O" in m.upper():
                return m
        for m in ordered:
            if "PROD" in m.upper():
                return m
        return ordered[0]

    h2 = pick_h2()
    co2 = pick_co2()
    out: Dict[str, str] = {}
    if h2 is not None:
        out["target"] = h2
    if co2 is not None:
        out["tailgas"] = co2
    return out


def resolve_gas_readout_nodes_from_process_fixed_map(
    process_id: str,
    node_names: List[str],
) -> Dict[str, str]:
    """Resolve fixed H2/CO2 readout node names for a known process id."""
    entry = PROCESS_FIXED_GAS_READOUT.get(process_id)
    if entry is None:
        return {}
    h2_name, co2_name = entry
    missing = [name for name in (h2_name, co2_name) if name not in node_names]
    if missing:
        raise ValueError(
            f"process_id={process_id!r} fixed gas mapping references unknown node(s): {missing}. "
            f"Known nodes: {node_names}"
        )
    return {
        "target": h2_name,
        "tailgas": co2_name,
    }


def parse_adjacency_sheet(ws: Worksheet) -> Tuple[List[str], List[List[int]]]:
    node_names_col = []
    for c in range(2, ws.max_column + 1):
        val = _clean_str(ws.cell(1, c).value)
        if val is not None:
            node_names_col.append(val)

    node_names_row = []
    for r in range(2, ws.max_row + 1):
        val = _clean_str(ws.cell(r, 1).value)
        if val is not None:
            node_names_row.append(val)

    if node_names_col != node_names_row:
        raise ValueError("Adjacency sheet row/column node names do not match.")

    node_names = node_names_col
    src_list: List[int] = []
    dst_list: List[int] = []

    n = len(node_names)
    for r in range(n):
        for c in range(n):
            v = ws.cell(r + 2, c + 2).value
            if v in (1, 1.0):
                src_list.append(r)
                dst_list.append(c)

    return node_names, [src_list, dst_list]


def _find_node_rows(ws: Worksheet) -> Dict[str, int]:
    row_map: Dict[str, int] = {}
    for r in range(3, ws.max_row + 1):
        node_name = _clean_str(ws.cell(r, 1).value)
        if node_name is not None:
            row_map[node_name] = r
    return row_map


def _find_present_slot_columns(ws: Worksheet) -> Dict[str, int]:
    slot_to_col: Dict[str, int] = {}
    for c in range(3, ws.max_column + 1):
        slot = _clean_str(ws.cell(2, c).value)
        if slot in OPER_FEATURE_SLOTS:
            slot_to_col[slot] = c
    return slot_to_col


def _find_bfp_readout_column(ws: Worksheet) -> Optional[int]:
    """Optional node-sheet column (header row 2) for fixed-task readout mapping."""
    for c in range(3, ws.max_column + 1):
        h = _clean_str(ws.cell(2, c).value)
        if h is None:
            continue
        key = h.upper().replace(" ", "_")
        if key in ("BFP_READOUT", "BFP_FIXED_READOUT"):
            return c
    return None


def parse_fixed_readout_from_node_sheet(
    ws: Worksheet,
    row_map: Dict[str, int],
    node_names: List[str],
    nodes: Dict[str, NodeSpec],
) -> Dict[str, str]:
    """
    Parse optional node-sheet column ``BFP_READOUT`` (header on row 2, same row as oper slot names).

    For each graph node row, the cell may list fixed supervised readout slots: ``target``, ``tailgas``,
    or both (separated by comma, semicolon, pipe, or whitespace). Only ``role=sink`` (OUT_*) nodes are
    allowed. Each slot must appear on exactly one node row per workbook.
    """
    col = _find_bfp_readout_column(ws)
    if col is None:
        return {}
    assignment: Dict[str, str] = {}
    valid_slots = frozenset({"target", "tailgas"})
    for node_name in node_names:
        r = row_map.get(node_name)
        if r is None:
            continue
        raw = ws.cell(r, col).value
        if raw is None:
            continue
        s = str(raw).strip()
        if not s:
            continue
        parts = [p.strip().lower() for p in re.split(r"[,;|/\s]+", s) if p.strip()]
        for p in parts:
            if p not in valid_slots:
                raise ValueError(
                    f"Invalid BFP_READOUT token {p!r} for node {node_name!r} on sheet {ws.title!r}. "
                    f"Allowed: target, tailgas (combine with comma, semicolon, or pipe)."
                )
            if p in assignment and assignment[p] != node_name:
                raise ValueError(
                    f"BFP_READOUT assigns {p!r} to both {assignment[p]!r} and {node_name!r}. "
                    f"Each slot must appear on exactly one node row."
                )
            node = nodes.get(node_name)
            if node is None:
                raise ValueError(f"BFP_READOUT references unknown node {node_name!r}.")
            if node.role != "sink":
                raise ValueError(
                    f"BFP_READOUT: node {node_name!r} for slot {p!r} must be role=sink (OUT_*), "
                    f"got role={node.role!r}."
                )
            assignment[p] = node_name
    return assignment


def parse_node_sheet(ws: Worksheet, node_names: List[str]) -> Dict[str, NodeSpec]:
    row_map = _find_node_rows(ws)
    slot_to_col = _find_present_slot_columns(ws)

    missing_nodes = [n for n in node_names if n not in row_map]
    if missing_nodes:
        raise ValueError(f"Node sheet is missing nodes: {missing_nodes}")

    nodes: Dict[str, NodeSpec] = {}

    for node_name in node_names:
        r = row_map[node_name]
        type_raw = _clean_str(ws.cell(r, 2).value)
        if type_raw is None:
            raise ValueError(f"Missing type for node {node_name}")

        role = infer_node_role(node_name)
        type_canonical = canonicalize_type(type_raw)
        hx_role = infer_hx_role(type_raw)

        features: Dict[str, CellSpec] = {}
        for slot in OPER_FEATURE_SLOTS:
            if slot not in slot_to_col:
                features[slot] = CellSpec(kind="missing", value=None)
                continue

            c = slot_to_col[slot]
            raw_value = ws.cell(r, c).value

            if type_canonical == "burner" and slot == "press":
                cell_spec = parse_burner_pressure_cell(raw_value)
            else:
                cell_spec = parse_cell_spec(raw_value)

            features[slot] = cell_spec

        nodes[node_name] = NodeSpec(
            name=node_name,
            role=role,
            type_raw=type_raw,
            type_canonical=type_canonical,
            hx_role=hx_role,
            features=features,
        )

    return nodes


def _find_decoder_header_row(ws: Worksheet) -> int:
    category_names = set(DECODER_CATEGORY_MAP.keys())
    for r in range(1, min(ws.max_row, 5) + 1):
        hits = 0
        for c in range(1, ws.max_column + 1):
            val = _clean_str(ws.cell(r, c).value)
            if val in category_names:
                hits += 1
        if hits >= 2:
            return r
    raise ValueError("Could not find decoder header row.")


def parse_decoder_sheet(ws: Worksheet, valid_node_names: Optional[set[str]] = None) -> Dict[str, Dict[str, str]]:
    header_row = _find_decoder_header_row(ws)

    col_to_category: Dict[int, str] = {}
    for c in range(2, ws.max_column + 1):
        raw_header = _clean_str(ws.cell(header_row, c).value)
        if raw_header in DECODER_CATEGORY_MAP:
            col_to_category[c] = DECODER_CATEGORY_MAP[raw_header]

    decoder_spec: Dict[str, Dict[str, str]] = {v: {} for v in DECODER_CATEGORY_MAP.values()}

    for r in range(header_row + 1, ws.max_row + 1):
        node_name = _clean_str(ws.cell(r, 1).value)
        if node_name is None:
            continue

        if valid_node_names is not None and node_name not in valid_node_names:
            raise ValueError(f"Decoder node '{node_name}' is not in graph node names.")

        for c, category in col_to_category.items():
            raw_target = _clean_str(ws.cell(r, c).value)
            if raw_target is not None:
                decoder_spec[category][node_name] = raw_target

    return decoder_spec


def parse_process_file(filepath: str, process_id: str) -> ProcessSpec:
    wb = load_workbook(filepath, data_only=False)
    sheet_roles = detect_sheet_roles(wb)

    node_names, edge_index = parse_adjacency_sheet(sheet_roles["adjacency"])
    node_ws = sheet_roles["node_spec"]
    row_map = _find_node_rows(node_ws)
    nodes = parse_node_sheet(node_ws, node_names)
    fixed_readout = parse_fixed_readout_from_node_sheet(node_ws, row_map, node_names, nodes)
    highlighted = collect_yellow_sink_names_second_sheet(wb)
    gas_readout = resolve_gas_readout_nodes_from_highlight_list(highlighted, node_names)
    if gas_readout:
        fixed_readout = {**fixed_readout, **gas_readout}
    process_fixed = resolve_gas_readout_nodes_from_process_fixed_map(process_id, node_names)
    if process_fixed:
        # Highest priority: explicit process-id mapping provided by project owner.
        fixed_readout = {**fixed_readout, **process_fixed}
    decoder_spec = parse_decoder_sheet(sheet_roles["decoder"], valid_node_names=set(node_names))

    return ProcessSpec(
        process_id=process_id,
        node_names=node_names,
        edge_index=edge_index,
        nodes=nodes,
        decoder_spec=decoder_spec,
        fixed_readout=fixed_readout,
    )